"""ppt-maker pptx export: slides measured in the page (export.js) -> an editable .pptx. python-pptx is imported lazily."""
import base64, binascii, io, os, re, secrets
from pathlib import Path

PX = 9525  # EMU per CSS px
MAX_IMG = 10 * 1024 * 1024
ALIGN = {"l": 1, "c": 2, "r": 3, "j": 4}  # PP_ALIGN values
SHAPES = {"rect": 1, "round": 5, "ellipse": 9, "triangle": 7}  # MSO_SHAPE values
DATA_URL = re.compile(r"data:image/(png|jpeg|gif|webp);base64,(.+)", re.S)


class Missing(RuntimeError):
    pass


def need_pptx():
    try:
        import pptx  # noqa: F401
    except ImportError as e:
        raise Missing("导出需要 python-pptx：pip install python-pptx") from e


def E(v):
    return int(round(float(v) * PX))


def image_bytes(src):
    m = DATA_URL.fullmatch(src or "")
    if not m:
        raise ValueError("不是图片")
    try:
        data = base64.b64decode(m.group(2), validate=False)
    except binascii.Error as e:
        raise ValueError("图片数据坏了") from e
    if len(data) > MAX_IMG:
        raise ValueError("图片超过 10MB")
    return data


def build(data, out):
    need_pptx()
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_CONNECTOR
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    def rgb(h):
        return RGBColor.from_string(str(h)[:6].upper())

    def alpha(fill, a):  # transparency on a solid fill
        if a is not None and a < 1:
            clr = fill._xPr.find(qn("a:solidFill"))[0]
            el = clr.makeelement(qn("a:alpha"), {"val": str(int(a * 100000))})
            clr.append(el)

    def plain(sh):  # drop the theme style: no blue fill, no shadow, black text
        st = sh._element.find(qn("p:style"))
        if st is not None:
            sh._element.remove(st)
        return sh

    def set_fill(fill, hexv, a=1):
        if hexv:
            fill.solid()
            fill.fore_color.rgb = rgb(hexv)
            alpha(fill, a)
        else:
            fill.background()

    def set_line(line, hexv, w, a=1, dash=None):  # <a:ln> children in schema order: fill, prstDash, …, headEnd, tailEnd
        if hexv and w:
            line.color.rgb = rgb(hexv)
            line.width = Emu(E(w))
            alpha(line.fill, a)
            if dash:
                ln = line._get_or_add_ln()
                ln.find(qn("a:solidFill")).addnext(ln.makeelement(qn("a:prstDash"), {"val": dash}))
        else:
            line.fill.background()

    def ends(line, start, end):  # arrowheads take the line's colour; size sm | med | lg
        ln = line._get_or_add_ln()
        for size, tag in ((start, "a:headEnd"), (end, "a:tailEnd")):
            if size:
                ln.append(ln.makeelement(qn(tag), {"type": "triangle", "w": size, "len": size}))

    def runs_into(para, p):
        para.alignment = ALIGN.get(p.get("align"))
        if p.get("lh"):
            para.line_spacing = Pt(float(p["lh"]) * 0.75)
        lvl = int(p.get("level") or 0)
        para.level = min(lvl, 8)
        if p.get("indent") and not p.get("bullet"):  # the first line starts after a leading badge or icon, as on the page
            para._p.get_or_add_pPr().set("indent", str(E(p["indent"])))
        if p.get("bullet"):
            pPr = para._p.get_or_add_pPr()
            pPr.set("marL", str(E(24 * (lvl + 1))))
            pPr.set("indent", str(-E(18)))
            if p["bullet"] == "•":
                pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))
            else:
                pPr.append(pPr.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"}))
        for r in p.get("runs") or []:
            if r.get("br"):
                para.add_line_break()
                continue
            run = para.add_run()
            run.text = r.get("t", "")
            f = run.font
            f.size = Pt(float(r.get("size") or 16) * 0.75)
            f.bold, f.italic, f.underline = bool(r.get("b")), bool(r.get("i")), bool(r.get("u"))
            if r.get("color"):
                f.color.rgb = rgb(r["color"])
                a = r.get("a")
                if a is not None and a < 1:  # semi-transparent text
                    clr = run._r.get_or_add_rPr().find(qn("a:solidFill"))[0]
                    clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(a * 100000))}))
            if r.get("font"):
                f.name = r["font"]
                rPr = run._r.get_or_add_rPr()
                rPr.find(qn("a:latin")).addnext(rPr.makeelement(qn("a:ea"), {"typeface": r.get("ea") or r["font"]}))
            if r.get("s"):
                run._r.get_or_add_rPr().set("strike", "sngStrike")

    def text_into(tf, paras, middle=False, wrap=True):
        tf.word_wrap = wrap
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        if middle:
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, p in enumerate(paras or []):
            runs_into(tf.paragraphs[0] if i == 0 else tf.add_paragraph(), p)

    def add_text(s, it):
        tb = s.shapes.add_textbox(E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"]))
        tb.rotation = float(it.get("rot") or 0)
        text_into(tb.text_frame, it["paras"], wrap=it.get("wrap", True) is not False)
        return "text"

    def corner(sh, it):  # a radius in px -> PowerPoint's rounding, a fraction of the shorter side
        if it.get("r"):
            sh.adjustments[0] = min(0.5, float(it["r"]) / max(1.0, min(float(it["w"]), float(it["h"]))))

    def add_box(s, it):
        kind = SHAPES["round"] if it.get("r") else SHAPES["rect"]
        sh = plain(s.shapes.add_shape(kind, E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"])))
        corner(sh, it)
        set_fill(sh.fill, it.get("fill"), it.get("fa", 1))
        set_line(sh.line, it.get("line"), it.get("lw"))
        return "shape"

    def add_shape(s, it):
        sh = plain(s.shapes.add_shape(SHAPES[it["shape"]], E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"])))
        sh.rotation = float(it.get("rot") or 0)
        if it["shape"] == "round":
            corner(sh, it)
        set_fill(sh.fill, it.get("fill"), it.get("fa", 1))
        set_line(sh.line, it.get("line"), it.get("lw"), it.get("la", 1), it.get("dash"))
        if it.get("paras"):
            text_into(sh.text_frame, it["paras"], middle=True)
        return "shape"

    def add_path(s, it):  # an SVG path (absolute canvas px, M/L/C/Z) as a freeform: colours, size and place stay editable
        segs = it.get("segs") or []
        pts = [(float(v[i]), float(v[i + 1])) for v in segs if v[0] in "MLC" for i in range(1, len(v), 2)]
        if len(pts) < 2:
            raise ValueError("没有可画的路径")
        x0, y0 = min(p[0] for p in pts), min(p[1] for p in pts)
        w, h = max(1.0, max(p[0] for p in pts) - x0), max(1.0, max(p[1] for p in pts) - y0)
        sh = plain(s.shapes.add_shape(SHAPES["rect"], E(x0), E(y0), E(w), E(h)))
        spPr = sh._element.spPr
        geom = spPr.makeelement(qn("a:custGeom"), {})
        for tag in ("a:avLst", "a:gdLst", "a:ahLst", "a:cxnLst"):
            geom.append(geom.makeelement(qn(tag), {}))
        geom.append(geom.makeelement(qn("a:rect"), {"l": "0", "t": "0", "r": "r", "b": "b"}))
        lst = geom.makeelement(qn("a:pathLst"), {})
        path = lst.makeelement(qn("a:path"), {"w": str(E(w)), "h": str(E(h))})
        if not it.get("fill") and not any(v[0] == "Z" for v in segs):
            path.set("fill", "none")  # an open line; a filled open path is filled as if closed, like SVG
        tags = {"M": "a:moveTo", "L": "a:lnTo", "C": "a:cubicBezTo", "Z": "a:close"}
        for v in segs:
            el = path.makeelement(qn(tags[v[0]]), {})
            for i in range(1, len(v), 2):
                el.append(el.makeelement(qn("a:pt"), {"x": str(E(float(v[i]) - x0)), "y": str(E(float(v[i + 1]) - y0))}))
            path.append(el)
        lst.append(path)
        geom.append(lst)
        spPr.replace(spPr.find(qn("a:prstGeom")), geom)
        set_fill(sh.fill, it.get("fill"), it.get("fa", 1))
        set_line(sh.line, it.get("line"), it.get("lw"), it.get("la", 1), it.get("dash"))
        if it.get("line") and it.get("lw"):
            ends(sh.line, it.get("start"), it.get("end"))
        return "shape"

    def add_group(s, it):  # one PowerPoint group per <svg>; a child that fails is skipped, the rest stay
        grp = s.shapes.add_group_shape()
        if not sum(place(grp, c) for c in it.get("items") or []):
            grp._element.getparent().remove(grp._element)

    def add_line(s, it):
        c = plain(s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(it["x1"]), E(it["y1"]), E(it["x2"]), E(it["y2"])))
        c.line.color.rgb = rgb(it.get("color") or "000000")
        c.line.width = Emu(E(it.get("lw") or 1))
        ln = c.line._get_or_add_ln()
        for key, tag in (("start", "a:headEnd"), ("end", "a:tailEnd")):
            if it.get(key):
                ln.append(ln.makeelement(qn(tag), {"type": "triangle"}))
        return "shape"

    def add_pic(s, it):
        pic = s.shapes.add_picture(io.BytesIO(image_bytes(it["src"])), E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"]))
        crop = it.get("crop") or {}
        pic.crop_left, pic.crop_top = float(crop.get("l", 0)), float(crop.get("t", 0))
        pic.crop_right, pic.crop_bottom = float(crop.get("r", 0)), float(crop.get("b", 0))
        return "image" if it["k"] == "image" else "raster"

    def add_table(s, it):
        rows, cols = len(it["cells"]), len(it["cells"][0])
        gf = s.shapes.add_table(rows, cols, E(it["x"]), E(it["y"]), E(it["w"]), E(it["h"]))
        t = gf.table
        t.first_row = t.horz_banding = False
        for j, w in enumerate(it.get("cols") or []):
            t.columns[j].width = Emu(E(w))
        for i, h in enumerate(it.get("rows") or []):
            t.rows[i].height = Emu(E(h))
        for i, row in enumerate(it["cells"]):
            for j, cd in enumerate(row):
                cell = t.cell(i, j)
                if it.get("border") and it.get("bw"):  # schema order: lnL, lnR, lnT, lnB, then the fill (PowerPoint repairs anything else)
                    tcPr = cell._tc.get_or_add_tcPr()
                    for k, side in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                        ln = tcPr.makeelement(qn(side), {"w": str(E(it["bw"]))})
                        sf = ln.makeelement(qn("a:solidFill"), {})
                        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": str(it["border"])[:6].upper()}))
                        ln.append(sf)
                        tcPr.insert(k, ln)
                set_fill(cell.fill, cd.get("fill"))
                cell.margin_left = cell.margin_right = Emu(E(4))
                cell.margin_top = cell.margin_bottom = Emu(E(2))
                text_into(cell.text_frame, cd.get("paras"))
        return "table"

    add = {"text": add_text, "box": add_box, "shape": add_shape, "line": add_line, "image": add_pic, "raster": add_pic, "table": add_table,
           "path": add_path, "group": add_group}
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1280 * PX), Emu(720 * PX)
    blank = prs.slide_layouts[6]
    rep = {"slides": 0, "text": 0, "shape": 0, "table": 0, "image": 0, "raster": [], "skipped": list(data.get("skipped") or [])}

    def place(container, it):  # the slide or a group: True when the item made it in
        tree = container.shapes._spTree
        n = len(tree)
        try:
            kind = add[it["k"]](container, it)
        except Exception as e:  # noqa: BLE001 — one bad item (e.g. PIL's OSError on a non-image) must not sink the deck
            for el in list(tree)[n:]:  # nor leave half of itself behind
                tree.remove(el)
            rep["skipped"].append({"id": it.get("id", ""), "why": f"无法生成：{e}"})
            return False
        if kind == "raster":
            rep["raster"].append({"id": it.get("id", ""), "why": it.get("why", "")})
        elif kind:  # a group counts its children, not itself
            rep[kind] += 1
        return len(tree) > n

    for sd in data["slides"]:
        s = prs.slides.add_slide(blank)
        rep["slides"] += 1
        if sd.get("bg"):
            s.background.fill.solid()
            s.background.fill.fore_color.rgb = rgb(sd["bg"])
        for it in sd.get("items") or []:
            place(s, it)
    out = Path(out)
    tmp = out.with_name(f".{out.stem}-{secrets.token_hex(4)}.part")  # never shared by two exports
    try:
        prs.save(str(tmp))
        os.replace(tmp, out)
    finally:
        tmp.unlink(missing_ok=True)
    return rep


def summary(rep):
    s = f"导出完成：{rep['slides']} 页，{rep['text']} 个文本框、{rep['shape']} 个形状、{rep['table']} 个表格、{rep['image']} 张图片"
    if rep.get("raster"):
        s += f"；{len(rep['raster'])} 处转成了图片（" + "；".join(f"{r['id']}：{r['why']}" for r in rep["raster"]) + "）"
    if rep.get("skipped"):
        s += f"；{len(rep['skipped'])} 处没能导出（" + "；".join(f"{r['id']}：{r['why']}" for r in rep["skipped"]) + "）"
    return s
