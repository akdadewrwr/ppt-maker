import base64, contextlib, http.client, io, json, os, shutil, socket, subprocess, sys, tempfile, threading, unittest, zipfile
import time as _time
from pathlib import Path
from unittest import mock
from urllib.parse import quote

sys.path.insert(0, str(Path(__file__).resolve().parent))
import server
import kb
import importer
import exporter

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "tools"))
import build_icons

TEMPLATE = Path(__file__).resolve().parent / "template.html"
SERVER_PY = str(Path(__file__).resolve().parent / "server.py")


class ServerCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.deck = self.tmp / "deck"
        self.deck.mkdir()
        shutil.copy(TEMPLATE, self.deck / "deck.html")
        self.srv = server.make_server(self.deck, port=0, token="tok")
        self.port = self.srv.server_address[1]
        threading.Thread(target=self.srv.serve_forever, daemon=True).start()

    def tearDown(self):
        self.srv.shutdown()
        self.srv.server_close()
        shutil.rmtree(self.tmp, ignore_errors=True)

    def req(self, method, path, body=None, host=None):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        data = json.dumps(body, ensure_ascii=False).encode() if body is not None else None
        c.request(method, path, data, {"Host": host or f"127.0.0.1:{self.port}", "Content-Type": "application/json"})
        r = c.getresponse()
        out = (r.status, r.read())
        c.close()
        return out


class PageTests(ServerCase):
    def test_index_requires_token(self):
        self.assertEqual(self.req("GET", "/")[0], 403)
        self.assertEqual(self.req("GET", "/?t=bad")[0], 403)
        status, body = self.req("GET", "/?t=tok")
        self.assertEqual(status, 200)
        html = body.decode()
        self.assertIn("window.PPT=", html)
        self.assertIn("/_ppt/editor.js", html)
        self.assertLess(html.index("/_ppt/editor.js"), html.lower().rindex("</body>"))

    def test_inject_without_body_tag(self):
        (self.deck / "deck.html").write_text("<html><SECTION class=slide id=s1></SECTION>", encoding="utf-8")
        self.assertIn("/_ppt/editor.js", self.req("GET", "/?t=tok")[1].decode())
        (self.deck / "deck.html").write_text("<html><body>x</BODY></html>", encoding="utf-8")
        html = self.req("GET", "/?t=tok")[1].decode()
        self.assertLess(html.index("/_ppt/editor.js"), html.index("</BODY>"))

    def test_nosniff(self):  # review M5: an uploaded "image" that is also HTML must never render as a page
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", "/_ppt/template.html", headers={"Host": f"127.0.0.1:{self.port}"})
        r = c.getresponse()
        r.read()
        c.close()
        self.assertEqual(r.getheader("X-Content-Type-Options"), "nosniff")

    def test_bad_host_rejected(self):
        self.assertEqual(self.req("GET", "/?t=tok", host=f"evil.com:{self.port}")[0], 403)
        self.assertEqual(self.req("GET", "/_ppt/server.py", host="evil.com")[0], 403)
        self.assertEqual(self.req("GET", "/?t=tok", host=f"localhost:{self.port}")[0], 200)

    def test_static_files_and_traversal(self):
        self.assertEqual(self.req("GET", "/_ppt/template.html")[0], 200)
        (self.deck / "assets").mkdir()
        (self.deck / "assets" / "a.png").write_bytes(b"\x89PNG")
        self.assertEqual(self.req("GET", "/assets/a.png")[0], 200)
        server.runtime(self.deck)
        (self.deck / ".ppt" / "inbox.jsonl").write_text("{}\n", encoding="utf-8")
        for bad in ("/.ppt/inbox.jsonl", "/_ppt/../../secret", "/%2e%2e/%2e%2e/x", "/_ppt/..%2f..%2fx", "/nope.png"):
            self.assertEqual(self.req("GET", bad)[0], 404, bad)

    def test_port_fallback(self):
        busy = socket.socket()
        busy.bind(("127.0.0.1", 0))
        busy.listen()
        try:
            srv = server.make_server(self.deck, port=busy.getsockname()[1])
            self.assertNotEqual(srv.server_address[1], busy.getsockname()[1])
            srv.server_close()
        finally:
            busy.close()


PNG = "data:image/png;base64," + __import__("base64").b64encode(b"\x89PNG\r\n\x1a\nfake").decode()


def inbox(deck):
    return server.read_jsonl(deck / ".ppt" / "inbox.jsonl")


class EventTests(ServerCase):
    def test_post_intent_with_shot(self):
        ev = {"type": "intent", "slide": 1, "slide_id": "s-cover", "rect": {"x": 0, "y": 0, "w": 80, "h": 80},
              "cells": "1-1/1-1", "blocks": ["b-cover-title"], "text": "标题改成蓝色渐变", "shot": PNG, "id": "hack"}
        status, body = self.req("POST", "/api/event?t=tok", ev)
        self.assertEqual(status, 200)
        self.assertEqual(json.loads(body)["id"], "e1")
        [stored] = inbox(self.deck)
        self.assertEqual(stored["id"], "e1")                      # client id ignored
        self.assertEqual(stored["text"], "标题改成蓝色渐变")         # Chinese round-trip
        self.assertEqual(stored["shot"], ".ppt/shots/e1.png")
        self.assertTrue((self.deck / ".ppt" / "shots" / "e1.png").read_bytes().startswith(b"\x89PNG"))
        self.assertEqual(json.loads(self.req("POST", "/api/event?t=tok", {"type": "answer", "re": "e1", "text": "好"})[1])["id"], "e2")

    def test_blank_area_selection_accepted(self):
        status, _ = self.req("POST", "/api/event?t=tok", {"type": "intent", "slide_id": "s-cover", "cells": "1-2/8-9", "blocks": [], "text": "这里加脚注"})
        self.assertEqual(status, 200)

    def test_post_requires_token(self):
        self.assertEqual(self.req("POST", "/api/event", {"type": "intent", "text": "x"})[0], 403)
        self.assertEqual(self.req("POST", "/api/event?t=no", {"type": "intent", "text": "x"})[0], 403)
        self.assertEqual(inbox(self.deck), [])

    def test_rejected_write_still_answers_403(self):
        # unread request bodies make Windows reset the socket instead of delivering the 403
        # repro: body arrives after the server already answered and closed -> RST, client never sees the status
        def slow(method, path):
            body = b'{"type": "intent", "text": "x"}'
            s = socket.create_connection(("127.0.0.1", self.port), timeout=5)
            s.sendall(f"{method} {path} HTTP/1.1\r\nHost: 127.0.0.1:{self.port}\r\n"
                      f"Content-Type: application/json\r\nContent-Length: {len(body)}\r\n\r\n".encode())
            _time.sleep(0.3)
            try:
                s.sendall(body)
                _time.sleep(0.3)
                return int(s.recv(64).split()[1])
            except OSError as e:
                return repr(e)
            finally:
                s.close()
        self.assertEqual(slow("POST", "/api/event"), 403)
        self.assertEqual(slow("PUT", "/api/deck?t=no"), 403)
        self.assertEqual(slow("POST", "/api/nope?t=tok"), 404)

    def test_post_rejects_bad_input(self):
        for bad in ({"type": "edit"}, {"type": "nope"}, [], {"type": "intent", "text": "字" * 4001},
                    {"type": "intent", "shot": "data:image/png;base64,R0lGODlh"}, {"type": "intent", "shot": "data:,!!!"}):
            self.assertEqual(self.req("POST", "/api/event?t=tok", bad)[0], 400, bad)
        self.assertEqual(inbox(self.deck), [])

    def test_put_deck_version_check(self):
        ver = json.loads(self.req("GET", "/api/state?t=tok")[1])["version"]
        html = (self.deck / "deck.html").read_text(encoding="utf-8").replace("副标题", "组会·第三周")
        status, body = self.req("PUT", "/api/deck?t=tok", {"version": ver, "html": html, "block": "b-cover-sub", "slide_id": "s-cover"})
        self.assertEqual(status, 200)
        new = json.loads(body)["version"]
        self.assertNotEqual(new, ver)
        self.assertIn("组会·第三周", (self.deck / "deck.html").read_text(encoding="utf-8"))
        self.assertEqual(inbox(self.deck)[-1]["type"], "edit")
        self.assertEqual(inbox(self.deck)[-1]["block"], "b-cover-sub")
        status, body = self.req("PUT", "/api/deck?t=tok", {"version": ver, "html": html})   # stale
        self.assertEqual(status, 409)
        self.assertEqual(json.loads(body)["version"], new)
        self.assertEqual(self.req("PUT", "/api/deck?t=tok", {"version": new, "html": "not html"})[0], 400)
        self.assertEqual(self.req("PUT", "/api/deck", {"version": new, "html": html})[0], 403)

    def test_state_statuses_and_online(self):
        self.req("POST", "/api/event?t=tok", {"type": "intent", "slide_id": "s-cover", "text": "a"})
        self.req("POST", "/api/event?t=tok", {"type": "intent", "slide_id": "s-cover", "text": "b"})
        self.req("POST", "/api/event?t=tok", {"type": "lint", "issues": []})
        out = server.runtime(self.deck) / "outbox.jsonl"
        server.append_jsonl(out, {"re": "e1", "text": "", "ack": True})
        server.append_jsonl(out, {"re": "e2", "text": "", "ack": True})
        server.append_jsonl(out, {"re": "e2", "text": "已改", "ack": False})
        s = json.loads(self.req("GET", "/api/state?t=tok")[1])
        self.assertFalse(s["online"])
        self.assertEqual([(e["id"], e["status"], e["reply"]) for e in s["events"]],
                         [("e1", "processing", None), ("e2", "done", "已改")])   # lint not listed
        (self.deck / ".ppt" / "heartbeat").write_text("x")
        self.assertTrue(json.loads(self.req("GET", "/api/state?t=tok")[1])["online"])



class ReviewFixTests(ServerCase):
    def test_wrong_field_types_rejected(self):  # I2
        for bad in ({"type": "intent", "blocks": [1]}, {"type": "intent", "text": 5},
                    {"type": "intent", "slide_id": ["x"]}, {"type": "lint", "issues": "x"}):
            self.assertEqual(self.req("POST", "/api/event?t=tok", bad)[0], 400, bad)

    def test_tail_survives_malformed_event(self):  # I2
        server.append_jsonl(server.runtime(self.deck) / "inbox.jsonl", {"id": "e1", "type": "intent", "blocks": [1], "text": 7})
        server.append_jsonl(server.runtime(self.deck) / "inbox.jsonl", {"id": "e2", "type": "intent", "text": "好"})
        lines = []
        server.tail(self.deck, lines.append, stop=lambda: True)
        self.assertEqual([l.split()[0] for l in lines], ["e1", "e2"])

    def test_read_jsonl_tolerates_torn_utf8(self):  # M1
        p = server.runtime(self.deck) / "inbox.jsonl"
        p.write_bytes('{"id": "e1", "type": "intent", "text": "好"}\n{"id": "e2", "text": "'.encode() + "中".encode()[:2])
        self.assertEqual([e["id"] for e in server.read_jsonl(p)], ["e1"])

    def test_token_persists_across_restarts(self):  # I2b
        a = server.make_server(self.deck, port=0)
        b = server.make_server(self.deck, port=0)
        try:
            self.assertEqual(a.token, b.token)
            self.assertEqual((self.deck / ".ppt" / "token").read_text(encoding="utf-8"), a.token)
        finally:
            a.server_close(); b.server_close()

    def test_get_raw_deck(self):  # I3: page rebuilds from the file text, not the live DOM
        status, body = self.req("GET", "/api/deck?t=tok")
        self.assertEqual(status, 200)
        d = json.loads(body)
        self.assertEqual(d["html"], (self.deck / "deck.html").read_text(encoding="utf-8"))
        self.assertEqual(d["version"], server.deck_version(self.deck))
        self.assertEqual(self.req("GET", "/api/deck")[0], 403)


class QuoteTests(ServerCase):
    def test_quote_intent_stored_without_shot(self):
        q = {"text": "副标题", "before": "", "after": ""}
        ev = {"type": "intent", "slide_id": "s-cover", "blocks": ["b-cover-sub"], "text": "更正式", "quote": q}
        self.assertEqual(self.req("POST", "/api/event?t=tok", ev)[0], 200)
        stored = inbox(self.deck)[-1]
        self.assertEqual(stored["quote"], q)
        self.assertNotIn("shot", stored)

    def test_bad_quote_rejected(self):
        for q in ("副标题", {"text": 5}, {"before": "x"}, {"text": "x" * 4001}):
            self.assertEqual(self.req("POST", "/api/event?t=tok", {"type": "intent", "quote": q})[0], 400, q)

    def test_summary_shows_quote(self):
        s = server.summary({"id": "e1", "type": "intent", "slide_id": "s1", "blocks": ["b"], "text": "改正式",
                            "quote": {"text": "在幻灯片上指给 agent 看，而不是说给它听"}})
        self.assertIn("quote:「在幻灯片上指给 agent 看，而不是", s)


PNG_BYTES = b"\x89PNG\r\n\x1a\n" + b"\0" * 32


class AssetTests(ServerCase):
    def post_raw(self, path, data, ctype="image/png"):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("POST", path, data, {"Host": f"127.0.0.1:{self.port}", "Content-Type": ctype})
        r = c.getresponse()
        out = (r.status, r.read())
        c.close()
        return out

    def test_png_saved_by_content_hash_and_reused(self):
        status, body = self.post_raw("/api/asset?t=tok", PNG_BYTES)
        self.assertEqual(status, 200)
        src = json.loads(body)["src"]
        self.assertRegex(src, r"^assets/[0-9a-f]{12}\.png$")
        self.assertEqual((self.deck / src).read_bytes(), PNG_BYTES)
        self.assertEqual(json.loads(self.post_raw("/api/asset?t=tok", PNG_BYTES)[1])["src"], src)
        self.assertEqual(len(list((self.deck / "assets").iterdir())), 1)
        self.assertEqual(self.req("GET", "/" + src)[0], 200)  # the page loads it like any deck file

    def test_type_from_magic_bytes_not_header(self):
        for data, ext in ((b"\xff\xd8\xff\xe0" + b"0" * 20, "jpg"), (b"GIF89a" + b"0" * 20, "gif"),
                          (b"RIFF\0\0\0\0WEBPVP8 " + b"0" * 20, "webp")):
            status, body = self.post_raw("/api/asset?t=tok", data, "application/octet-stream")
            self.assertEqual((status, json.loads(body)["src"].rsplit(".", 1)[1]), (200, ext))

    def test_svg_and_junk_rejected(self):
        for data in (b'<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>', b"hello", b""):
            self.assertEqual(self.post_raw("/api/asset?t=tok", data, "image/svg+xml")[0], 415)
        self.assertFalse((self.deck / "assets").exists())

    def test_needs_token(self):
        self.assertEqual(self.post_raw("/api/asset", PNG_BYTES)[0], 403)
        self.assertEqual(self.post_raw("/api/asset?t=bad", PNG_BYTES)[0], 403)

    def test_too_large(self):
        with mock.patch.object(server, "MAX_BODY", 16):
            self.assertEqual(self.post_raw("/api/asset?t=tok", PNG_BYTES)[0], 413)
        self.assertFalse((self.deck / "assets").exists())


class NewDeckTests(unittest.TestCase):  # a new deck comes from the template, made by the program, inside the project only
    def test_missing_deck_is_made_from_the_template_in_place(self):
        proj = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, proj, True)
        deck = proj / "decks" / "试用"
        self.assertTrue(server.ensure_deck(deck))
        tpl = (Path(server.__file__).parent / "template.html").read_text(encoding="utf-8")
        self.assertEqual((deck / "deck.html").read_text(encoding="utf-8"), tpl)
        self.assertEqual(sorted(p.relative_to(proj).as_posix() for p in proj.rglob("*")), ["decks", "decks/试用", "decks/试用/deck.html"])

    def test_starting_the_server_on_a_missing_deck_makes_it(self):
        proj = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, proj, True)
        deck = proj / "decks" / "新稿"
        out = io.StringIO()
        with mock.patch.object(server, "make_server", side_effect=SystemExit("stop")), contextlib.redirect_stdout(out), self.assertRaises(SystemExit):
            server.main(["--no-open", str(deck)])
        self.assertTrue((deck / "deck.html").is_file())
        self.assertIn("已从模板新建", out.getvalue())
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):  # other commands still need a real deck
            server.main(["--tail", str(proj / "decks" / "没有")])
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            server.main(["--stop", str(proj / "decks" / "没有")])
        self.assertFalse((proj / "decks" / "没有").exists())

    def test_an_existing_deck_is_left_alone(self):
        proj = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, proj, True)
        (proj / "deck.html").write_text("<p>mine</p>", encoding="utf-8")
        self.assertFalse(server.ensure_deck(proj))
        self.assertEqual((proj / "deck.html").read_text(encoding="utf-8"), "<p>mine</p>")


class DetachTests(unittest.TestCase):  # the server runs on its own, not as a Claude Code background task
    def setUp(self):
        self.proj = Path(tempfile.mkdtemp())
        self.deck = self.proj / "decks" / "d"

    def tearDown(self):
        subprocess.run([sys.executable, SERVER_PY, "--stop", str(self.deck)], capture_output=True, timeout=30)
        shutil.rmtree(self.proj, ignore_errors=True)

    def run_cli(self, *args):
        return subprocess.run([sys.executable, SERVER_PY, *args], capture_output=True, text=True, encoding="utf-8", timeout=60)

    def get(self, info, path):
        c = http.client.HTTPConnection("127.0.0.1", info["port"], timeout=5)
        c.request("GET", path, headers={"Host": f"127.0.0.1:{info['port']}"})
        status = c.getresponse().status
        c.close()
        return status

    def test_detach_outlives_the_launcher_is_reused_and_stops(self):
        r1 = self.run_cli("--detach", "--no-open", "--port", "8840", str(self.deck))  # returns: the launcher has exited
        self.assertEqual(r1.returncode, 0, r1.stderr)
        self.assertIn("已从模板新建", r1.stdout)
        self.assertIn("新启动", r1.stdout)
        info = json.loads((server.runtime(self.deck) / "server.json").read_text(encoding="utf-8"))
        self.assertIn(info["url"], r1.stdout)
        self.assertEqual(self.get(info, f"/?t={info['token']}"), 200)
        r2 = self.run_cli("--detach", "--no-open", "--port", "8840", str(self.deck))
        self.assertIn("已在运行", r2.stdout)
        self.assertIn(info["url"], r2.stdout)
        self.assertEqual(json.loads((server.runtime(self.deck) / "server.json").read_text(encoding="utf-8"))["pid"], info["pid"])  # not a second server
        r3 = self.run_cli("--stop", str(self.deck))
        self.assertIn("已停止", r3.stdout)
        self.assertFalse((server.runtime(self.deck) / "server.json").exists())
        with self.assertRaises(OSError):
            self.get(info, "/")

    def test_an_idle_server_exits_on_its_own(self):
        server.ensure_deck(self.deck)
        srv = server.make_server(self.deck, port=0, token="tok", idle=0.5)
        t = threading.Thread(target=srv.serve_forever, daemon=True)
        t.start()
        t.join(5)
        srv.server_close()
        self.assertFalse(t.is_alive())

    def test_tail_tells_recent_work_and_a_stopped_server(self):
        server.ensure_deck(self.deck)
        rt = server.runtime(self.deck)
        for i in range(1, 13):
            server.append_jsonl(rt / "inbox.jsonl", {"id": f"e{i}", "type": "intent", "text": f"改第{i}处"})
            if i < 12:
                server.reply(self.deck, f"e{i}", f"改好了{i}")
        (rt / "server.json").write_text(json.dumps({"pid": 0, "port": 9, "url": "x", "token": "t"}), encoding="utf-8")  # a server that is gone
        lines = []
        server.tail(self.deck, lines.append, stop=lambda: True)
        recent = [l for l in lines if l.startswith("最近完成")]
        self.assertEqual(len(recent), 10)
        self.assertIn("e11", recent[-1])
        self.assertIn("改好了11", recent[-1])
        self.assertTrue(any(l.startswith("e12 ") for l in lines))
        self.assertTrue(any("服务已停止" in l for l in lines))


class WaitTests(unittest.TestCase):  # --wait: Codex's waiting loop
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.deck = self.tmp / "deck"
        server.ensure_deck(self.deck)
        self.inbox = server.runtime(self.deck) / "inbox.jsonl"

    def tearDown(self):
        shutil.rmtree(self.tmp, ignore_errors=True)

    def add(self, eid, typ="intent", **k):
        server.append_jsonl(self.inbox, {"id": eid, "type": typ, "text": f"做{eid}", **k})

    def wait(self, rounds=40, **k):  # stop after `rounds` polls so a broken wait can never hang the suite
        n, lines = [0], []
        def stop():
            n[0] += 1
            return n[0] > rounds
        got = server.wait_events(self.deck, lines.append, stop=stop, interval=0.05, **k)
        return got, lines

    def test_wait_ignores_edits_then_returns_with_them(self):
        self.add("e1", "edit", label="移动 b-left")
        threading.Timer(0.3, lambda: self.add("e2")).start()
        got, lines = self.wait()
        self.assertEqual(got, "events")
        self.assertEqual([l.split()[0] for l in lines], ["e1", "e2"])

    def test_wait_never_repeats_acked(self):
        self.add("e1")
        self.assertEqual(self.wait()[0], "events")
        server.reply(self.deck, "e1", ack=True)
        got, lines = self.wait(rounds=3)
        self.assertNotEqual(got, "events")
        self.assertFalse(any(l.startswith("e1 ") for l in lines))

    def test_wait_returns_unacked_before_cursor(self):
        self.add("e1")
        self.wait()
        got, lines = self.wait(rounds=3)
        self.assertEqual(got, "events")
        self.assertTrue(lines[0].startswith("e1 "))

    def test_wait_recent_lists_done_and_processing(self):
        for i in range(1, 13):
            self.add(f"e{i}")
            if i < 12:
                server.reply(self.deck, f"e{i}", f"改好了{i}")
        server.reply(self.deck, "e12", ack=True)
        got, lines = self.wait(rounds=3, recent=True)
        self.assertEqual(got, "events")
        self.assertEqual(len([l for l in lines if l.startswith("最近完成")]), 10)
        self.assertTrue(any(l.startswith("e12 ") for l in lines))

    def test_wait_idle_stop(self):
        t = _time.time()
        got, lines = self.wait(rounds=1000, idle=1)
        self.assertEqual(got, "idle")
        self.assertLess(_time.time() - t, 3)
        self.assertIn("已停止监听", lines[-1])
        self.assertIn("--wait", lines[-1])
        self.assertTrue(server.idle_stop_text(900).startswith("15 分钟没有新操作"))

    def test_wait_edit_resets_idle(self):
        threading.Timer(0.8, lambda: self.add("e1", "edit", label="改字")).start()
        t = _time.time()
        got, _ = self.wait(rounds=1000, idle=1.5)
        self.assertEqual(got, "idle")
        self.assertGreater(_time.time() - t, 2)

    def test_wait_returns_when_server_stops(self):
        (server.runtime(self.deck) / "server.json").write_text(json.dumps({"pid": 0, "port": 9, "url": "x", "token": "t"}), encoding="utf-8")
        got, lines = self.wait(rounds=1000, idle=60)
        self.assertEqual(got, "stopped")
        self.assertTrue(any("服务已停止" in l for l in lines))

    def test_wait_notices_a_server_that_stopped_cleanly(self):  # review I1: --stop and a clean exit delete server.json
        ups = iter([True])
        with mock.patch.object(server, "server_up", side_effect=lambda d: next(ups, None)), mock.patch.object(server, "SERVER_CHECK_SECS", 0):
            got, lines = self.wait(rounds=1000, idle=60)
        self.assertEqual(got, "stopped")
        self.assertTrue(any("服务已停止" in l for l in lines))

    def test_wait_cli(self):
        self.add("e1")
        out = io.StringIO()
        with contextlib.redirect_stdout(out):
            server.main([str(self.deck), "--wait", "--idle", "1"])
        self.assertIn("e1 intent", out.getvalue())


class AgentTests(ServerCase):  # which agent is listening, shown on the page and in history labels
    def state(self):
        return json.loads(self.req("GET", "/api/state?t=tok")[1])

    def test_agent_defaults_to_claude(self):
        self.assertEqual(self.state()["agent"], "Claude")
        self.assertEqual(server.claude_label(self.deck), "Claude 修改")

    def test_agent_name_from_listener(self):
        server.write_agent(self.deck, "Codex")
        self.assertEqual(self.state()["agent"], "Codex")
        server.reply(self.deck, "e3", ack=True)
        self.assertEqual(server.claude_label(self.deck), "Codex 处理 #e3")

    def test_agent_name_is_sanitised(self):
        f = server.runtime(self.deck) / "agent"
        f.write_text("  \n", encoding="utf-8")
        self.assertEqual(server.read_agent(self.deck), "Claude")
        f.write_text("A" * 50 + "\nx", encoding="utf-8")
        self.assertEqual(server.read_agent(self.deck), "A" * 20)

    def test_online_while_processing(self):
        hb = server.runtime(self.deck) / "heartbeat"
        hb.write_text("x")
        old = lambda secs: os.utime(hb, (_time.time() - secs, _time.time() - secs))
        old(120)
        self.assertFalse(self.state()["online"])  # nothing in progress
        self.req("POST", "/api/event?t=tok", {"type": "intent", "slide_id": "s-cover", "text": "a"})
        server.reply(self.deck, "e1", ack=True)
        old(120)
        self.assertTrue(self.state()["online"])
        old(700)
        self.assertFalse(self.state()["online"])

    def test_reply_and_ack_touch_heartbeat(self):
        hb = server.runtime(self.deck) / "heartbeat"
        self.assertFalse(hb.exists())
        server.reply(self.deck, "e1", ack=True)
        self.assertTrue(hb.exists())
        self.assertLess(_time.time() - hb.stat().st_mtime, 5)

    def test_tail_writes_agent(self):
        server.tail(self.deck, lambda s: None, stop=lambda: True, agent="Codex")
        self.assertEqual(server.read_agent(self.deck), "Codex")


class PromptTests(unittest.TestCase):  # one set of rules for every agent, an entry per agent
    ROOT = Path(__file__).resolve().parent.parent

    def read(self, rel):
        return (self.ROOT / rel).read_text(encoding="utf-8")

    def test_common_rules_name_no_claude_tools(self):
        common = self.read("prompts/ppt-common.md")
        self.assertIn("## 处理事件", common)
        for word in ("Monitor", "WebFetch", "run_in_background", "Claude Code", "用 Read"):
            self.assertNotIn(word, common)

    def test_both_entries_read_the_common_rules(self):
        claude, codex = self.read(".claude/commands/ppt.md"), self.read("codex/skills/ppt/SKILL.md")
        self.assertIn("prompts/ppt-common.md", claude)
        self.assertIn("prompts/ppt-common.md", codex)
        self.assertTrue(codex.startswith("---\nname: ppt\n"))
        self.assertIn("\ndescription:", codex.split("---")[1])
        for s in ("--wait --recent --agent Codex", "--wait --agent Codex", "已停止监听", "windows.sandbox", "unelevated", "{{PPT_HOME}}"):
            self.assertIn(s, codex)

    def test_page_text_names_no_fixed_agent(self):  # review I2: every user-facing "Claude" goes through PPTAgent.name()
        import re
        fixed = []
        for f in sorted((self.ROOT / "ppt").glob("*.js")):
            for n, line in enumerate(f.read_text(encoding="utf-8").splitlines(), 1):
                code = line.split(" // ")[0]
                if not code.lstrip().startswith("//") and re.search(r"(['\"`])[^'\"`]*Claude[^'\"`]*\1", code):
                    fixed.append(f"{f.name}:{n}")
        self.assertEqual(fixed, [])

    def test_codex_acks_everything_it_is_shown(self):  # review: an event never acked makes every --wait return at once
        self.assertIn("每个输出的事件都要", self.read("codex/skills/ppt/SKILL.md"))

    def test_claude_entry_keeps_its_loop(self):
        claude = self.read(".claude/commands/ppt.md")
        for s in ("--detach", "--tail", "Monitor"):
            self.assertIn(s, claude)
        self.assertNotIn("## deck.html 格式规则", claude)  # moved to the common rules

    def test_install_codex(self):
        home = Path(tempfile.mkdtemp())
        try:
            d = server.install_codex(home)
            self.assertEqual(d, home / ".agents" / "skills" / "ppt")
            text = (d / "SKILL.md").read_text(encoding="utf-8")
            self.assertNotIn("{{PPT_HOME}}", text)
            self.assertIn(self.ROOT.as_posix() + "/ppt/server.py", text)
            self.assertTrue((home / ".ppt-maker").is_dir())  # an existing dir, so it can be a Codex writable root
            server.install_codex(home)  # again: replaces, no error
            self.assertEqual((d / "SKILL.md").read_text(encoding="utf-8"), text)
        finally:
            shutil.rmtree(home, ignore_errors=True)

    def test_install_codex_cli(self):
        home = Path(tempfile.mkdtemp())
        try:
            out = io.StringIO()
            with mock.patch.object(Path, "home", return_value=home), contextlib.redirect_stdout(out):
                server.main(["--install-codex"])
            self.assertIn('sandbox = "unelevated"', out.getvalue())
            self.assertIn("$ppt", out.getvalue())
            self.assertTrue((home / ".agents" / "skills" / "ppt" / "SKILL.md").is_file())
        finally:
            shutil.rmtree(home, ignore_errors=True)

    def test_install_claude_cli(self):
        home = Path(tempfile.mkdtemp())
        try:
            out = io.StringIO()
            with mock.patch.object(Path, "home", return_value=home), contextlib.redirect_stdout(out):
                server.main(["--install-claude"])
            self.assertIn("/ppt", out.getvalue())
            text = (home / ".claude" / "commands" / "ppt.md").read_text(encoding="utf-8")
            self.assertNotIn("{{PPT_HOME}}", text)
            self.assertIn(self.ROOT.as_posix() + "/.claude/commands/ppt.md", text)  # the global entry defers to the repo's
            self.assertIn(self.ROOT.as_posix() + "/ppt/server.py", text)
            self.assertTrue((home / ".ppt-maker").is_dir())
            with mock.patch.object(Path, "home", return_value=home), contextlib.redirect_stdout(out):
                server.main(["--install-claude"])  # again: replaces, no error
        finally:
            shutil.rmtree(home, ignore_errors=True)


PNG_1PX = base64.b64encode(bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000154a24f5d0000000049454e44ae426082")).decode()


def sample_export():
    run = lambda t, **k: {"t": t, "font": "Microsoft YaHei", "size": 40, "color": "112233", "a": 1, "b": False, "i": False, "u": False, "s": False, **k}
    para = lambda *runs, **k: {"align": "l", "lh": None, "bullet": "", "level": 0, "runs": list(runs), **k}
    return {"slides": [
        {"id": "s-cover", "bg": "0f172a", "items": [
            {"k": "box", "x": 80, "y": 80, "w": 400, "h": 200, "fill": "1e293b", "fa": 1, "line": None, "lw": 0, "r": 12},
            {"k": "text", "x": 100, "y": 100, "w": 600, "h": 60, "paras": [para(run("季度复盘", b=True), {"br": True}, run("第二行"), align="c")]},
            {"k": "text", "x": 100, "y": 300, "w": 600, "h": 90, "paras": [para(run("要点一"), bullet="•"), para(run("要点二"), bullet="•", level=1)]},
            {"k": "shape", "shape": "ellipse", "x": 700, "y": 100, "w": 200, "h": 100, "rot": 30, "fill": "38bdf8", "fa": 1, "line": "ffffff", "lw": 2,
             "paras": [para(run("圆", size=20), align="c")]},
            {"k": "line", "x1": 100, "y1": 600, "x2": 500, "y2": 600, "color": "fbbf24", "lw": 3, "start": False, "end": True},
        ]},
        {"id": "s-data", "bg": None, "items": [
            {"k": "table", "x": 80, "y": 80, "w": 400, "h": 80, "cols": [200, 200], "rows": [40, 40], "border": "cccccc", "bw": 1,
             "cells": [[{"paras": [para(run("指标", size=16))], "fill": "eeeeee"}, {"paras": [para(run("数值", size=16))], "fill": None}],
                       [{"paras": [para(run("用户", size=16))], "fill": None}, {"paras": [para(run("820 万", size=16))], "fill": None}]]},
            {"k": "image", "x": 600, "y": 80, "w": 100, "h": 100, "src": "data:image/png;base64," + PNG_1PX, "crop": {"l": 0.1, "t": 0, "r": 0.1, "b": 0}},
            {"k": "raster", "x": 600, "y": 300, "w": 100, "h": 100, "src": "data:image/png;base64," + PNG_1PX, "id": "b-chart", "why": "SVG 插图"},
        ]}], "skipped": [{"id": "b-ext", "why": "图片取不到"}]}


class ExporterTests(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, True)

    def test_build_makes_editable_native_objects(self):
        from pptx import Presentation
        from pptx.util import Emu
        out = self.tmp / "out" / "d.pptx"
        out.parent.mkdir()
        rep = exporter.build(sample_export(), out)
        self.assertEqual({k: rep[k] for k in ("slides", "text", "shape", "table", "image")}, {"slides": 2, "text": 2, "shape": 3, "table": 1, "image": 1})  # box + ellipse + line
        self.assertEqual(rep["raster"], [{"id": "b-chart", "why": "SVG 插图"}])
        self.assertEqual(rep["skipped"], [{"id": "b-ext", "why": "图片取不到"}])
        self.assertEqual([p.name for p in out.parent.iterdir()], ["d.pptx"])  # no temp file left
        prs = Presentation(str(out))
        self.assertEqual((prs.slide_width, prs.slide_height), (Emu(1280 * 9525), Emu(720 * 9525)))
        s1, s2 = prs.slides
        self.assertEqual(str(s1.background.fill.fore_color.rgb), "0F172A")
        tb = [sh for sh in s1.shapes if sh.has_text_frame and sh.text_frame.text.startswith("季度复盘")][0]
        self.assertEqual((tb.left, tb.top, tb.width), (Emu(100 * 9525), Emu(100 * 9525), Emu(600 * 9525)))
        r0 = tb.text_frame.paragraphs[0].runs[0]
        self.assertEqual((r0.font.size.pt, r0.font.bold, str(r0.font.color.rgb), r0.font.name), (30.0, True, "112233", "Microsoft YaHei"))
        self.assertIn("\v", tb.text_frame.paragraphs[0].text)  # the <br> is a line break inside the paragraph
        self.assertIn('typeface="Microsoft YaHei"', r0._r.xml.split("<a:ea")[1][:60])  # east asian font set too
        bul = [sh for sh in s1.shapes if sh.has_text_frame and sh.text_frame.text.startswith("要点一")][0].text_frame.paragraphs
        self.assertIn("buChar", bul[0]._p.xml)
        self.assertEqual(bul[1].level, 1)
        ell = [sh for sh in s1.shapes if sh.has_text_frame and sh.text_frame.text == "圆"][0]
        from pptx.enum.shapes import MSO_SHAPE
        self.assertEqual((ell.auto_shape_type, ell.rotation), (MSO_SHAPE.OVAL, 30.0))
        self.assertNotIn("<p:style>", ell._element.xml)  # no theme style: no surprise blue fill or shadow
        line = [sh for sh in s1.shapes if sh.shape_type == 9 or sh.__class__.__name__ == "Connector"][0]
        self.assertIn("tailEnd", line._element.xml)
        self.assertNotIn("headEnd", line._element.xml)
        tbl = [sh for sh in s2.shapes if sh.has_table][0].table
        self.assertEqual([[c.text for c in r.cells] for r in tbl.rows], [["指标", "数值"], ["用户", "820 万"]])
        pics = [sh for sh in s2.shapes if sh.shape_type == 13]
        self.assertEqual(len(pics), 2)
        self.assertAlmostEqual(pics[0].crop_left, 0.1, places=3)

    def test_one_line_text_never_wraps(self):  # acceptance: PowerPoint's wider metrics wrapped a one-line subtitle
        from pptx import Presentation
        para = {"align": "l", "lh": None, "bullet": "", "level": 0, "runs": [{"t": "产品组 · 2026 年 10 月 · 内部", "size": 28}]}
        data = {"slides": [{"id": "s", "bg": None, "items": [{"k": "text", "x": 0, "y": 0, "w": 300, "h": 40, "wrap": False, "paras": [para]},
                                                            {"k": "text", "x": 0, "y": 100, "w": 300, "h": 80, "paras": [para]}]}]}
        exporter.build(data, self.tmp / "w.pptx")
        a, b = Presentation(str(self.tmp / "w.pptx")).slides[0].shapes
        self.assertEqual((a.text_frame.word_wrap, b.text_frame.word_wrap), (False, True))

    def test_table_border_xml_is_in_schema_order(self):  # review C1: PowerPoint repairs lnL..lnB placed around the fill
        from pptx import Presentation
        from pptx.oxml.ns import qn
        exporter.build(sample_export(), self.tmp / "t.pptx")
        tbl = [sh for sh in Presentation(str(self.tmp / "t.pptx")).slides[1].shapes if sh.has_table][0].table
        for cell in (tbl.cell(0, 0), tbl.cell(0, 1)):
            tags = [c.tag.split("}")[1] for c in cell._tc.tcPr]
            self.assertEqual(tags[:4], ["lnL", "lnR", "lnT", "lnB"], tags)
            self.assertIn(tags[4], ("solidFill", "noFill"), tags)

    def test_a_corrupt_image_is_skipped_not_fatal(self):  # review I3: PIL raises OSError for bytes that are not an image
        bad = "data:image/png;base64," + base64.b64encode(b"<html>CDN error</html>").decode()
        data = {"slides": [{"id": "s", "bg": None, "items": [{"k": "image", "id": "b-pic", "x": 0, "y": 0, "w": 9, "h": 9, "src": bad, "crop": None},
                                                            {"k": "text", "x": 0, "y": 20, "w": 90, "h": 20, "paras": [{"runs": [{"t": "还在", "size": 16}]}]}]}]}
        rep = exporter.build(data, self.tmp / "c.pptx")
        self.assertEqual((rep["text"], [s["id"] for s in rep["skipped"]]), (1, ["b-pic"]))

    def test_east_asian_font_is_its_own(self):  # review: spec 2.2 — Chinese text gets a CJK face, not the Latin one
        para = {"runs": [{"t": "中文 abc", "size": 16, "font": "Inter", "ea": "Microsoft YaHei"}]}
        exporter.build({"slides": [{"id": "s", "bg": None, "items": [{"k": "text", "x": 0, "y": 0, "w": 90, "h": 20, "paras": [para]}]}]}, self.tmp / "f.pptx")
        from pptx import Presentation
        xml = Presentation(str(self.tmp / "f.pptx")).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]._r.xml
        self.assertIn('<a:latin typeface="Inter"', xml)
        self.assertIn('<a:ea typeface="Microsoft YaHei"', xml)

    def test_a_failing_item_leaves_no_half_built_shape(self):  # minor: a ragged table failed after its frame was added
        from pptx import Presentation
        cell = lambda t: {"paras": [{"runs": [{"t": t, "size": 16}]}], "fill": None}
        data = {"slides": [{"id": "s", "bg": None, "items": [{"k": "table", "id": "b-t", "x": 0, "y": 0, "w": 200, "h": 80, "cols": [100, 100], "rows": [40, 40],
                                                            "border": None, "bw": 0, "cells": [[cell("a"), cell("b")], [cell("c"), cell("d"), cell("e")]]}]}]}
        rep = exporter.build(data, self.tmp / "h.pptx")
        self.assertEqual([s["id"] for s in rep["skipped"]], ["b-t"])
        self.assertEqual(len(Presentation(str(self.tmp / "h.pptx")).slides[0].shapes), 0)

    def test_text_transparency_is_kept(self):  # minor: run alpha was ignored
        from pptx import Presentation
        para = {"runs": [{"t": "半透明", "size": 16, "color": "112233", "a": 0.5}]}
        exporter.build({"slides": [{"id": "s", "bg": None, "items": [{"k": "text", "x": 0, "y": 0, "w": 90, "h": 20, "paras": [para]}]}]}, self.tmp / "a.pptx")
        xml = Presentation(str(self.tmp / "a.pptx")).slides[0].shapes[0].text_frame.paragraphs[0].runs[0]._r.xml
        self.assertIn('<a:alpha val="50000"/>', xml)

    def test_summary_names_what_became_pictures(self):
        rep = {"slides": 4, "text": 18, "shape": 6, "table": 1, "image": 2, "raster": [{"id": "b-chart", "why": "SVG 插图"}], "skipped": []}
        self.assertEqual(exporter.summary(rep), "导出完成：4 页，18 个文本框、6 个形状、1 个表格、2 张图片；1 处转成了图片（b-chart：SVG 插图）")
        self.assertEqual(exporter.summary({**rep, "raster": [], "skipped": [{"id": "x", "why": "坏"}]}),
                         "导出完成：4 页，18 个文本框、6 个形状、1 个表格、2 张图片；1 处没能导出（x：坏）")

    def test_a_bad_item_is_skipped_not_fatal(self):
        data = {"slides": [{"id": "s", "bg": None, "items": [{"k": "image", "x": 0, "y": 0, "w": 9, "h": 9, "src": "data:text/html;base64,PGI+", "crop": None},
                                                            {"k": "nonsense"}]}]}
        rep = exporter.build(data, self.tmp / "x.pptx")
        self.assertEqual([s["why"][:5] for s in rep["skipped"]], ["无法生成：", "无法生成："])

    def test_missing_python_pptx_is_reported(self):
        with mock.patch.dict(sys.modules, {"pptx": None}):
            with self.assertRaises(exporter.Missing) as c:
                exporter.build(sample_export(), self.tmp / "y.pptx")
        self.assertEqual(str(c.exception), "导出需要 python-pptx：pip install python-pptx")


class ExportApiTests(ServerCase):
    def post_raw(self, path, data):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=30)
        c.request("POST", path, data, {"Host": f"127.0.0.1:{self.port}", "Content-Type": "application/json"})
        r = c.getresponse()
        out = (r.status, json.loads(r.read() or b"null"))
        c.close()
        return out

    def test_export_writes_a_file_and_offers_it_as_a_download(self):
        st, d = self.post_raw("/api/export?t=tok", json.dumps(sample_export()).encode())
        self.assertEqual(st, 200, d)
        f = Path(d["file"])
        self.assertEqual(f.parent, (self.deck / "export").resolve())
        self.assertRegex(f.name, r"^deck-\d{8}-\d{6}\.pptx$")
        self.assertTrue(d["summary"].startswith("导出完成：2 页"))
        self.assertEqual(json.loads((self.deck / ".ppt" / "export.json").read_text(encoding="utf-8"))["file"], d["file"])
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", d["url"], headers={"Host": f"127.0.0.1:{self.port}"})
        r = c.getresponse()
        self.assertEqual(r.status, 200)
        self.assertEqual(r.getheader("Content-Disposition"), "attachment; filename*=UTF-8''" + quote(f.name))
        self.assertEqual(r.read()[:2], b"PK")
        c.close()

    def test_a_failed_export_says_why_in_export_json(self):  # review I7: the terminal export reads it instead of waiting 120 s
        with mock.patch.object(exporter, "build", side_effect=exporter.Missing("导出需要 python-pptx：pip install python-pptx")):
            self.post_raw("/api/export?t=tok", json.dumps(sample_export()).encode())
        self.assertEqual(json.loads((self.deck / ".ppt" / "export.json").read_text(encoding="utf-8")), {"error": "导出需要 python-pptx：pip install python-pptx"})

    def test_two_exports_in_the_same_second_get_two_files(self):  # minor
        fixed = server.datetime(2026, 10, 16, 9, 30, 0)
        with mock.patch.object(server, "datetime", mock.Mock(now=mock.Mock(return_value=fixed))):
            a = self.post_raw("/api/export?t=tok", json.dumps(sample_export()).encode())[1]
            b = self.post_raw("/api/export?t=tok", json.dumps(sample_export()).encode())[1]
        self.assertNotEqual(a["file"], b["file"])
        self.assertTrue(Path(a["file"]).is_file() and Path(b["file"]).is_file())

    def test_export_needs_the_token_and_a_sane_body(self):
        self.assertEqual(self.post_raw("/api/export?t=bad", b"{}")[0], 403)
        self.assertEqual(self.post_raw("/api/export?t=tok", b"not json")[0], 400)
        self.assertEqual(self.post_raw("/api/export?t=tok", json.dumps({"slides": "x"}).encode())[0], 400)

    def test_export_over_50mb_is_refused(self):  # the limit is patched small; refused uploads are still read off so the 413 arrives
        with mock.patch.object(server, "MAX_EXPORT", 16):
            self.assertEqual(self.post_raw("/api/export?t=tok", b"x" * 1000)[0], 413)
        self.assertEqual(server.MAX_EXPORT, 50 * 1024 * 1024)

    def test_export_without_python_pptx_says_how_to_install(self):
        with mock.patch.object(exporter, "build", side_effect=exporter.Missing("导出需要 python-pptx：pip install python-pptx")):
            st, d = self.post_raw("/api/export?t=tok", json.dumps(sample_export()).encode())
        self.assertEqual((st, d["error"]), (503, "导出需要 python-pptx：pip install python-pptx"))

    def test_a_chinese_deck_name_downloads_with_its_name(self):
        (self.deck / "export").mkdir()
        (self.deck / "export" / "季度 复盘-1.pptx").write_bytes(b"PK")
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", "/export/" + quote("季度 复盘-1.pptx"), headers={"Host": f"127.0.0.1:{self.port}"})
        self.assertEqual(c.getresponse().getheader("Content-Disposition"), "attachment; filename*=UTF-8''" + quote("季度 复盘-1.pptx"))
        c.close()


@unittest.skipUnless(server.browser(), "needs Edge or Chrome")
class ExportCliTests(unittest.TestCase):
    def test_export_from_the_terminal_reports_a_failure_quickly(self):  # review I7
        proj = Path(tempfile.mkdtemp())
        deck = proj / "decks" / "坏"
        self.addCleanup(shutil.rmtree, proj, True)
        self.addCleanup(subprocess.run, [sys.executable, SERVER_PY, "--stop", str(deck)], capture_output=True, timeout=30)
        fake = proj / "fake"
        (fake / "pptx").mkdir(parents=True)
        (fake / "pptx" / "__init__.py").write_text("raise ImportError('no pptx here')", encoding="utf-8")
        server.ensure_deck(deck)
        env = {**os.environ, "PYTHONIOENCODING": "utf-8", "PYTHONPATH": str(fake)}  # the detached server inherits it
        t0 = _time.time()
        r = subprocess.run([sys.executable, SERVER_PY, "--export", "--port", "8865", str(deck)], capture_output=True, text=True, encoding="utf-8", timeout=200, env=env)
        self.assertNotEqual(r.returncode, 0)
        self.assertIn("导出需要 python-pptx", r.stdout + r.stderr)
        self.assertLess(_time.time() - t0, 60)

    def test_export_from_the_terminal(self):
        proj = Path(tempfile.mkdtemp())
        deck = proj / "decks" / "导出"
        self.addCleanup(shutil.rmtree, proj, True)
        self.addCleanup(subprocess.run, [sys.executable, SERVER_PY, "--stop", str(deck)], capture_output=True, timeout=30)
        env = {**os.environ, "PYTHONIOENCODING": "utf-8"}
        miss = subprocess.run([sys.executable, SERVER_PY, "--export", str(deck)], capture_output=True, text=True, encoding="utf-8", timeout=60, env=env)
        self.assertNotEqual(miss.returncode, 0)  # nothing to export: no deck is made for it
        self.assertFalse(deck.exists())
        server.ensure_deck(deck)
        r = subprocess.run([sys.executable, SERVER_PY, "--export", "--port", "8860", str(deck)], capture_output=True, text=True, encoding="utf-8", timeout=200, env=env)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        files = list((deck / "export").glob("*.pptx"))
        self.assertEqual(len(files), 1)
        self.assertIn(str(files[0].resolve()), r.stdout)
        self.assertIn("导出完成：1 页", r.stdout)


def make_pptx(path, width_in=13.333, height_in=7.5):
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width_in), Inches(height_in)
    s = prs.slides.add_slide(prs.slide_layouts[1])  # title + content placeholders: positions come from the layout
    s.shapes.title.text = "季度复盘"
    tb = s.shapes.add_textbox(Inches(1), Inches(5), Inches(4), Inches(1))
    r = tb.text_frame.paragraphs[0].add_run()
    r.text, r.font.size, r.font.bold = "大字", Pt(30), True
    r.font.color.rgb = RGBColor(0x11, 0x22, 0x33)
    sh = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(6), Inches(1), Inches(2), Inches(1))
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor(0x38, 0xBD, 0xF8)
    grp = s.shapes.add_group_shape()
    grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(9), Inches(4), Inches(1), Inches(1))
    s.shapes.add_picture(io.BytesIO(base64.b64decode(PNG_1PX)), Inches(10), Inches(1), Inches(1), Inches(1))
    t = s.shapes.add_table(2, 2, Inches(1), Inches(6), Inches(4), Inches(1)).table
    t.cell(0, 0).text = "指标"
    prs.slides.add_slide(prs.slide_layouts[6])
    prs.save(str(path))


class PptxImportTests(ServerCase):
    def up(self, data, name, mode="append"):
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=120)
        c.request("POST", f"/api/import?t=tok&mode={mode}&name={quote(name)}", data, {"Host": f"127.0.0.1:{self.port}"})
        r = c.getresponse()
        out = (r.status, json.loads(r.read() or b"null"))
        c.close()
        return out

    def test_pptx_is_taken_apart_for_claude(self):
        src = self.tmp / "old.pptx"
        make_pptx(src)
        with mock.patch.object(importer, "render_pdf", return_value=None):
            st, d = self.up(src.read_bytes(), "旧稿.pptx")
        self.assertEqual(st, 200, d)
        e = inbox(self.deck)[-1]
        self.assertEqual((e["type"], e["hint"], e["pdf"]), ("import", "pptx", None))
        sj = json.loads(Path(e["source"]).read_text(encoding="utf-8"))
        self.assertEqual((sj["size"]["w"], round(sj["size"]["h"]), sj["size"]["ratio"]), (1280, 720, "16:9"))
        self.assertEqual(len(sj["slides"]), 2)
        items = sj["slides"][0]["items"]
        title = [i for i in items if i["kind"] == "text" and i["paras"] and i["paras"][0]["runs"] and i["paras"][0]["runs"][0]["t"] == "季度复盘"][0]
        self.assertGreater(title["w"], 0)  # inherited from the layout placeholder
        big = [i for i in items if i["kind"] == "text" and "大字" in json.dumps(i, ensure_ascii=False)][0]
        self.assertEqual((round(big["x"]), round(big["y"])), (96, 480))
        run = big["paras"][0]["runs"][0]
        self.assertEqual((round(run["size"]), run["b"], run["color"]), (40, True, "112233"))
        oval = [i for i in items if i["kind"] == "shape" and i.get("shape") == "OVAL"][0]
        self.assertEqual(oval["fill"], "38BDF8")
        grp_rect = [i for i in items if i["kind"] == "shape" and i.get("shape") == "RECTANGLE"][0]
        self.assertEqual((round(grp_rect["x"]), round(grp_rect["y"])), (864, 384))  # group child in page coordinates
        pic = [i for i in items if i["kind"] == "picture"][0]
        self.assertTrue((Path(e["dir"]) / pic["file"]).is_file())
        tbl = [i for i in items if i["kind"] == "table"][0]
        self.assertEqual(tbl["cells"][0][0], "指标")
        self.assertRegex(sj["theme"]["colors"]["dk1"], r"^[0-9A-F]{6}$")  # a system colour gives its real value, not "windowText"
        self.assertRegex(sj["theme"]["colors"]["accent1"], r"^[0-9A-F]{6}$")

    def test_a_4_3_deck_keeps_its_ratio(self):
        src = self.tmp / "old43.pptx"
        make_pptx(src, 10, 7.5)
        with mock.patch.object(importer, "render_pdf", return_value=None):
            self.assertEqual(self.up(src.read_bytes(), "a.pptx")[0], 200)
        sj = json.loads(Path(inbox(self.deck)[-1]["source"]).read_text(encoding="utf-8"))
        self.assertEqual((sj["size"]["w"], sj["size"]["h"], sj["size"]["ratio"]), (1280, 960, "4:3"))

    def test_broken_pptx_and_old_ppt_are_refused_cleanly(self):
        self.assertEqual(self.up(b"PK\x03\x04 not really", "bad.pptx"), (400, {"error": "无法读取这个 pptx"}))
        self.assertEqual(self.up(b"\xd0\xcf\x11\xe0", "old.ppt")[1]["error"], "请先在 PowerPoint 里另存为 .pptx")
        self.assertEqual([p for p in (self.deck / "imports").iterdir()] if (self.deck / "imports").exists() else [], [])

    def test_odd_real_world_shapes_do_not_crash_the_import(self):  # review I4
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.oxml.ns import qn
        from pptx.util import Inches
        prs = Presentation()
        s = prs.slides.add_slide(prs.slide_layouts[6])
        grp = s.shapes.add_group_shape()
        grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(1), Inches(1), Inches(1))
        xf = grp._element.grpSpPr.find(qn("a:xfrm"))
        for t in ("a:chOff", "a:chExt"):  # optional in the schema
            xf.remove(xf.find(qn(t)))
        odd = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(3), Inches(1), Inches(1), Inches(1))
        pr = odd._element.spPr
        pr.remove(pr.find(qn("a:prstGeom")))  # no geometry: python-pptx can't name its type
        pic = s.shapes.add_picture(io.BytesIO(base64.b64decode(PNG_1PX)), Inches(5), Inches(1), Inches(1), Inches(1))
        s.part.related_part(pic._element.blipFill.blip.rEmbed)._blob = b"<svg xmlns='http://www.w3.org/2000/svg'/>"  # PIL can't read it
        src = self.tmp / "odd.pptx"
        prs.save(str(src))
        with mock.patch.object(importer, "render_pdf", return_value=None):
            st, d = self.up(src.read_bytes(), "odd.pptx")
        self.assertEqual(st, 200, d)
        items = json.loads(Path(inbox(self.deck)[-1]["source"]).read_text(encoding="utf-8"))["slides"][0]["items"]
        self.assertEqual(len(items), 3, items)
        rect = [i for i in items if i.get("shape") == "RECTANGLE"][0]
        self.assertEqual((round(rect["x"]), round(rect["y"])), (128, 128))  # default 10in-wide page: 1in = 128px

    def test_any_unexpected_import_error_is_a_clean_400(self):  # review I4: never a dropped connection
        src = self.tmp / "old.pptx"
        make_pptx(src)
        with mock.patch.object(importer, "pptx_extract", side_effect=RuntimeError("boom")):
            st, d = self.up(src.read_bytes(), "x.pptx")
        self.assertEqual((st, d["error"]), (400, "无法读取这个 pptx"))
        self.assertEqual(list((self.deck / "imports").iterdir()), [])

    def test_smartart_rotation_outlines_and_shared_pictures(self):  # minors: SmartArt type, group rotation, outlines, duplicate pictures
        from pptx import Presentation
        from pptx.chart.data import CategoryChartData
        from pptx.dml.color import RGBColor
        from pptx.enum.chart import XL_CHART_TYPE
        from pptx.enum.shapes import MSO_SHAPE
        from pptx.oxml.ns import qn
        from pptx.util import Inches, Pt
        prs = Presentation()
        prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
        s = prs.slides.add_slide(prs.slide_layouts[6])
        cd = CategoryChartData(); cd.categories = ["a"]; cd.add_series("s", (1,))
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1), Inches(2), Inches(2), cd)
        gf._element.find(".//" + qn("a:graphicData")).set("uri", "http://schemas.openxmlformats.org/drawingml/2006/diagram")
        grp = s.shapes.add_group_shape()
        grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(4), Inches(1), Inches(2), Inches(2))  # group spans 4..8in x, 1..3in y
        grp.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(7), Inches(2), Inches(1), Inches(1))  # off-centre child: (7.5, 2.5)in centre
        grp.rotation = 90
        ov = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(9), Inches(1), Inches(1), Inches(1))
        ov.line.color.rgb = RGBColor(0xFF, 0x00, 0x00); ov.line.width = Pt(3)
        for x in (10, 11):
            s.shapes.add_picture(io.BytesIO(base64.b64decode(PNG_1PX)), Inches(x), Inches(4), Inches(1), Inches(1))
        src = self.tmp / "m.pptx"
        prs.save(str(src))
        with mock.patch.object(importer, "render_pdf", return_value=None):
            self.assertEqual(self.up(src.read_bytes(), "m.pptx")[0], 200)
        e = inbox(self.deck)[-1]
        items = json.loads(Path(e["source"]).read_text(encoding="utf-8"))["slides"][0]["items"]
        self.assertEqual([i["type"] for i in items if i["kind"] == "unsupported"], ["smartart"])
        small = [i for i in items if i.get("shape") == "RECTANGLE" and round(i["w"]) == 96][0]
        # group centre (6in, 2in) = (576, 192)px; the child centre (720, 240)px turned 90° clockwise lands at (528, 336)
        self.assertEqual((round(small["x"] + small["w"] / 2), round(small["y"] + small["h"] / 2), small["rotation"]), (528, 336, 90.0))
        oval = [i for i in items if i.get("shape") == "OVAL"][0]
        self.assertEqual((oval["line"], round(oval["lw"])), ("FF0000", 4))
        pics = [i for i in items if i["kind"] == "picture"]
        self.assertEqual(len({p["file"] for p in pics}), 1)
        self.assertEqual(len(list(Path(e["media"]).iterdir())), 1)

    def test_a_pptx_that_unpacks_huge_is_refused(self):  # minor: no unpacked-size cap
        src = self.tmp / "big.pptx"
        make_pptx(src)
        with mock.patch.object(importer, "MAX_UNPACKED", 1000):
            self.assertEqual(self.up(src.read_bytes(), "big.pptx"), (400, {"error": "pptx 解压后超过 200MB"}))
        self.assertEqual(list((self.deck / "imports").iterdir()), [])

    def test_kill_tree_ends_grandchildren(self):  # minor: a timed-out LibreOffice / Edge left helper processes behind
        pidfile = self.tmp / "pid"
        code = ("import subprocess, sys, time; c = subprocess.Popen([sys.executable, '-c', 'import time; time.sleep(60)']);"
                f" open({str(pidfile)!r}, 'w').write(str(c.pid)); time.sleep(60)")
        p = subprocess.Popen([sys.executable, "-c", code])
        for _ in range(100):
            if pidfile.exists() and pidfile.read_text():
                break
            _time.sleep(0.1)
        child = int(pidfile.read_text())
        importer.kill_tree(p)
        _time.sleep(1)
        out = subprocess.run(["tasklist", "/FI", f"PID eq {child}"], capture_output=True, text=True).stdout if os.name == "nt" else ""
        self.assertNotIn(str(child), out)
        self.assertIsNotNone(p.poll())

    def test_no_python_pptx_gives_503(self):
        src = self.tmp / "old.pptx"
        make_pptx(src)
        with mock.patch.object(importer, "pptx_extract", side_effect=exporter.Missing("导出需要 python-pptx：pip install python-pptx")):
            self.assertEqual(self.up(src.read_bytes(), "x.pptx")[0], 503)
        self.assertEqual(list((self.deck / "imports").iterdir()), [])

    def test_render_pdf_without_libreoffice_is_none(self):
        with mock.patch.object(importer, "soffice", return_value=None):
            self.assertIsNone(importer.render_pdf(self.tmp / "x.pptx", self.tmp))

    def test_pdf_original_is_not_sandboxed(self):  # Chrome's PDF viewer won't open under CSP sandbox
        (self.deck / "imports" / "i-x").mkdir(parents=True)
        (self.deck / "imports" / "i-x" / "slides.pdf").write_bytes(b"%PDF-1.4")
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", "/imports/i-x/slides.pdf", headers={"Host": f"127.0.0.1:{self.port}"})
        self.assertIsNone(c.getresponse().getheader("Content-Security-Policy"))
        c.close()


class DrainTests(unittest.TestCase):
    def test_drain_discards_in_chunks(self):  # review M4: an oversized body must not be buffered whole
        reads = []

        class F(io.BytesIO):
            def read(self, n=-1):
                reads.append(n)
                return super().read(n)

        f = F(b"x" * 200_000)
        server.drain(f, 200_000)
        self.assertEqual((f.tell(), max(reads)), (200_000, 65536))


class SketchTests(ServerCase):
    def test_notes_kept_and_summarised(self):
        self.req("POST", "/api/event?t=tok", {"type": "intent", "kind": "sketch", "slide_id": "s-cover", "blocks": [],
                                              "text": "按草图修改这一页", "notes": ["标题", "这里放图表"]})
        e = inbox(self.deck)[-1]
        self.assertEqual(e["notes"], ["标题", "这里放图表"])
        self.assertIn("notes:「标题」「这里放图表」", server.summary(e))

    def test_bad_notes_rejected(self):
        for notes in (["ok", 5], ["x"] * 51, ["长" * 201], "标题"):
            self.assertEqual(self.req("POST", "/api/event?t=tok", {"type": "intent", "kind": "sketch", "text": "x", "notes": notes})[0], 400, notes)

    def test_unencodable_text_rejected_not_dropped(self):  # review I5: a lone surrogate jammed the page's event queue
        for body in (b'{"type": "intent", "text": "x", "notes": ["\\ud83d"]}', b'{"type": "intent", "text": "x\\ud83d"}'):
            c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
            c.request("POST", "/api/event?t=tok", body, {"Host": f"127.0.0.1:{self.port}", "Content-Type": "application/json"})
            r = c.getresponse()
            r.read()
            c.close()
            self.assertEqual(r.status, 400, body)

    def test_sketch_js_injected_before_toolbar(self):
        html = self.req("GET", "/?t=tok")[1].decode()
        at = [html.index(f"/_ppt/{n}") for n in ("icons-zh.js", "sketch.js", "toolbar.js")]
        self.assertEqual(at, sorted(at))


class DrawIconTests(ServerCase):
    def test_kind_kept_and_summarised(self):
        self.req("POST", "/api/event?t=tok", {"type": "intent", "kind": "draw-icon", "slide_id": "s-cover",
                                              "rect": {"x": 608, "y": 328, "w": 64, "h": 64}, "blocks": [], "text": "一只猫"})
        e = inbox(self.deck)[-1]
        self.assertEqual(e["kind"], "draw-icon")
        self.assertIn("intent kind:draw-icon slide:s-cover", server.summary(e))
        self.assertEqual(self.req("POST", "/api/event?t=tok", {"type": "intent", "kind": 5, "text": "x"})[0], 400)

    def test_icons_injected_between_ops_and_toolbar(self):
        html = self.req("GET", "/?t=tok")[1].decode()
        at = [html.index(f"/_ppt/{n}") for n in ("ops.js", "icons.js", "icons-zh.js", "toolbar.js", "editor.js")]
        self.assertEqual(at, sorted(at))

    def test_icons_js_optional(self):
        with mock.patch.object(server, "HERE", self.tmp):  # a checkout where build_icons.py was never run
            html = self.req("GET", "/?t=tok")[1].decode()
        self.assertNotIn("/_ppt/icons.js", html)
        self.assertIn("/_ppt/icons-zh.js", html)

    def test_template_has_object_rules(self):
        self.assertIn(".obj { position:absolute; box-sizing:border-box; }", TEMPLATE.read_text(encoding="utf-8"))


class HistoryTests(ServerCase):
    def put(self, html, label="改背景色"):
        ver = json.loads(self.req("GET", "/api/state?t=tok")[1])["version"]
        return self.req("PUT", "/api/deck?t=tok", {"version": ver, "html": html, "label": label})

    def deck_text(self):
        return (self.deck / "deck.html").read_text(encoding="utf-8")

    def test_claude_label_survives_a_fast_reply(self):  # Claude may reply before the 0.5 s watcher tick
        self.req("POST", "/api/event?t=tok", {"type": "intent", "slide_id": "s-cover", "text": "x"})
        server.reply(self.deck, "e1", ack=True)
        with server.LOCK:  # hold the watcher off so write + reply both land before the snapshot
            (self.deck / "deck.html").write_text(self.deck_text().replace("副标题", "CLAUDE"), encoding="utf-8")
            server.reply(self.deck, "e1", "改好了")
            self.srv.history.sync()
        self.assertEqual(self.srv.history.rows[-1]["label"], "Claude 处理 #e1")

    def test_put_records_user_snapshot(self):
        a = self.deck_text()
        self.assertEqual(self.put(a.replace("副标题", "B"))[0], 200)
        h = self.srv.history
        self.assertEqual((h.rows[-1]["by"], h.rows[-1]["label"]), ("user", "改背景色"))
        s = json.loads(self.req("GET", "/api/state?t=tok")[1])["history"]
        self.assertEqual((s["canUndo"], s["undoLabel"], s["canRedo"]), (True, "改背景色", False))
        self.assertEqual(inbox(self.deck)[-1]["label"], "改背景色")
        self.assertIn("改背景色", server.summary(inbox(self.deck)[-1]))

    def test_external_write_recorded_as_claude(self):
        server.reply(self.deck, "e3", ack=True)
        (self.deck / "deck.html").write_text(self.deck_text().replace("副标题", "C"), encoding="utf-8")
        _time.sleep(1.2)  # watcher runs every 0.5 s
        self.assertEqual((self.srv.history.rows[-1]["by"], self.srv.history.rows[-1]["label"]), ("claude", "Claude 处理 #e3"))

    def test_undo_redo_roundtrip(self):
        a = self.deck_text()
        self.put(a.replace("副标题", "B"))
        b = self.deck_text()
        self.assertEqual(self.req("POST", "/api/undo?t=tok", {})[0], 200)
        self.assertEqual(self.deck_text(), a)
        ev = inbox(self.deck)[-1]
        self.assertEqual((ev["type"], ev["op"], ev["label"]), ("edit", "undo", "撤销了 改背景色"))
        self.assertEqual(self.req("POST", "/api/redo?t=tok", {})[0], 200)
        self.assertEqual(self.deck_text(), b)
        self.req("POST", "/api/undo?t=tok", {})
        self.put(a.replace("副标题", "C"), "改字")          # a new change drops the redo branch
        self.assertEqual(self.req("POST", "/api/redo?t=tok", {})[0], 409)

    def test_nothing_to_undo_and_token(self):
        self.assertEqual(self.req("POST", "/api/undo?t=tok", {})[0], 409)
        self.assertEqual(self.req("POST", "/api/undo", {})[0], 403)
        self.assertEqual(self.req("POST", "/api/redo?t=no", {})[0], 403)

    def test_undo_not_rerecorded_by_watcher(self):
        self.put(self.deck_text().replace("副标题", "B"))
        self.req("POST", "/api/undo?t=tok", {})
        n = len(self.srv.history.rows)
        _time.sleep(1.2)
        self.assertEqual(len(self.srv.history.rows), n)

    def test_cap(self):
        old = server.HISTORY_MAX
        server.HISTORY_MAX = 5
        try:
            for i in range(8):
                self.put(self.deck_text().replace("</body>", f"<!--{i}--></body>"), f"第{i}步")
            h = self.srv.history
            self.assertEqual(len(h.rows), 5)
            self.assertEqual(len(list(h.dir.glob("*.html"))), 5)
            self.assertEqual(h.rows[-1]["label"], "第7步")
        finally:
            server.HISTORY_MAX = old

    def test_snapshot_failure_does_not_block_save(self):
        self.srv.history.dir = self.deck / "missing" / "dir"
        self.assertEqual(self.put(self.deck_text().replace("副标题", "B"))[0], 200)
        self.assertIn(">B<", self.deck_text())
        s = json.loads(self.req("GET", "/api/state?t=tok")[1])["history"]
        self.assertEqual((s["ok"], s["canUndo"]), (False, False))

    def test_undo_right_after_claude_write_keeps_it_recoverable(self):  # review C1
        a = self.deck_text()
        self.put(a.replace("副标题", "B"))
        b = self.deck_text()
        (self.deck / "deck.html").write_text(b.replace(">B<", ">CLAUDE<"), encoding="utf-8")
        self.assertEqual(self.req("POST", "/api/undo?t=tok", {})[0], 200)  # before the 0.5 s watcher tick
        self.assertEqual(self.deck_text(), b)
        self.assertEqual(inbox(self.deck)[-1]["label"], "撤销了 Claude 修改")
        self.req("POST", "/api/redo?t=tok", {})
        self.assertIn(">CLAUDE<", self.deck_text())

    def test_put_right_after_claude_write_gets_its_own_row(self):  # review C1
        (self.deck / "deck.html").write_text(self.deck_text().replace("副标题", "CLAUDE"), encoding="utf-8")
        self.put(self.deck_text().replace("CLAUDE", "USER"), "改字")
        rows = self.srv.history.rows
        self.assertEqual([(r["by"], r["label"]) for r in rows[-2:]], [("claude", "Claude 修改"), ("user", "改字")])

    def test_restart_keeps_position(self):
        a = self.deck_text()
        self.put(a.replace("副标题", "B"))
        self.req("POST", "/api/undo?t=tok", {})
        with server.LOCK:
            h2 = server.History(self.deck)
        self.assertEqual((h2.cur, len(h2.rows)), (self.srv.history.cur, len(self.srv.history.rows)))
        self.assertTrue(h2.info()["canRedo"])


class BuildIconsTests(unittest.TestCase):
    def test_build_from_package(self):
        tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, tmp, True)
        pkg = tmp / "package"
        (pkg / "icons").mkdir(parents=True)
        (pkg / "icons" / "house.svg").write_text(
            '<!-- @license lucide-static v9.9.9 - ISC -->\n<svg\n  class="lucide lucide-house"\n  viewBox="0 0 24 24"\n>\n'
            '  <path d="M3 10a2 2 0 0 1 1 1" />\n  <circle cx="12" cy="12" r="3" />\n</svg>\n', encoding="utf-8")
        (pkg / "icons" / "no-tags.svg").write_text('<svg viewBox="0 0 24 24"><path d="M1 1h2"/></svg>', encoding="utf-8")
        (pkg / "tags.json").write_text('{"house": ["home", "living"]}', encoding="utf-8")
        (pkg / "package.json").write_text('{"version": "9.9.9"}', encoding="utf-8")
        (pkg / "LICENSE").write_text("ISC License", encoding="utf-8")
        self.assertEqual(build_icons.build(pkg, tmp / "icons.js", tmp / "LICENSE-lucide"), 2)
        js = (tmp / "icons.js").read_text(encoding="utf-8")
        self.assertIn("lucide-static 9.9.9", js.splitlines()[0])
        data = json.loads(js.split("window.PPTIcons = ", 1)[1].strip().rstrip(";"))
        self.assertEqual(data["house"], {"tags": "home living", "body": '<path d="M3 10a2 2 0 0 1 1 1"/><circle cx="12" cy="12" r="3"/>'})
        self.assertEqual(data["no-tags"]["tags"], "")
        self.assertEqual((tmp / "LICENSE-lucide").read_text(encoding="utf-8"), "ISC License")


def zip_bytes(files):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in files.items():
            z.writestr(name, data)
    return buf.getvalue()


W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
A_NS = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
S_NS = 'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"'
DOCX = zip_bytes({"word/document.xml": f'<w:document {W_NS}><w:body><w:p><w:r><w:t>营收</w:t></w:r><w:r><w:t>12.3 亿</w:t></w:r></w:p>'
                                       f'<w:p><w:r><w:t>增长 18%</w:t></w:r></w:p></w:body></w:document>'})
_slide = lambda t: f'<p:sld xmlns:p="urn:p" {A_NS}><a:p><a:r><a:t>{t}</a:t></a:r></a:p></p:sld>'
PPTX = zip_bytes({"ppt/slides/slide2.xml": _slide("第二页"), "ppt/slides/slide10.xml": _slide("第十页"), "ppt/slides/slide1.xml": _slide("封面")})
XLSX = zip_bytes({"xl/workbook.xml": f'<workbook {S_NS}><sheets><sheet name="数据"/></sheets></workbook>',
                  "xl/sharedStrings.xml": f'<sst {S_NS}><si><t>地区</t></si><si><t>华东</t></si></sst>',
                  "xl/worksheets/sheet1.xml": f'<worksheet {S_NS}><sheetData><row><c t="s"><v>0</v></c><c><v>2024</v></c></row>'
                                              f'<row><c t="s"><v>1</v></c><c><v>12.3</v></c></row></sheetData></worksheet>'})


class KbCase(unittest.TestCase):
    def setUp(self):
        self.tmp = Path(tempfile.mkdtemp())
        self.addCleanup(shutil.rmtree, self.tmp, True)
        self.proj, self.home = self.tmp / "proj", self.tmp / "home"
        self.deck = self.proj / "decks" / "d1"
        self.deck.mkdir(parents=True)
        shutil.copy(TEMPLATE, self.deck / "deck.html")
        env = mock.patch.dict(os.environ, {"PPT_PROJECT_DIR": str(self.proj), "PPT_HOME_DIR": str(self.home)})
        env.start()
        self.addCleanup(env.stop)
        self.r = kb.root(self.deck, "project")


class KbModuleTests(KbCase):
    def test_roots_follow_the_overrides(self):
        self.assertEqual(kb.root(self.deck, "project"), self.proj / "kb")
        self.assertEqual(kb.root(self.deck, "global"), self.home / ".ppt-maker" / "kb")
        self.assertEqual(kb.skills_root(self.deck, "global"), self.home / ".claude" / "skills")
        with self.assertRaises(ValueError):
            kb.root(self.deck, "elsewhere")

    def test_extract_office(self):
        for name, data, want in (("a.docx", DOCX, "营收12.3 亿\n增长 18%"),
                                 ("b.pptx", PPTX, "## 第 1 页\n封面\n\n## 第 2 页\n第二页\n\n## 第 3 页\n第十页"),
                                 ("c.xlsx", XLSX, "## 工作表 数据\n地区\t2024\n华东\t12.3")):
            (self.tmp / name).write_bytes(data)
            self.assertEqual(kb.extract(self.tmp / name), want)

    def test_add_file_extracts_and_dedupes(self):
        item, created = kb.add_file(self.r, io.BytesIO(DOCX), len(DOCX), "报告 2025.docx", "file")
        self.assertTrue(created)
        self.assertRegex(item["id"], r"^k-[0-9a-z]{6}$")
        self.assertEqual((item["type"], item["title"], item["status"]), ("file", "报告 2025", "pending"))
        self.assertEqual((self.r / item["source"]).read_bytes(), DOCX)
        self.assertEqual((self.r / item["text"]).read_text(encoding="utf-8"), "营收12.3 亿\n增长 18%")
        again, created2 = kb.add_file(self.r, io.BytesIO(DOCX), len(DOCX), "另一个名字.docx", "file")
        self.assertEqual((again["id"], created2), (item["id"], False))
        self.assertEqual(len(kb.load(self.r)), 1)
        self.assertEqual(len(list((self.r / "files").glob("*.docx"))), 1)

    def test_add_file_validation(self):
        for name, typ, kind in (("x.exe", "file", None), ("a.png", "ref", None), ("a.png", "file", "visual"), ("a.png", "web", None)):
            with self.assertRaises(ValueError):
                kb.add_file(self.r, io.BytesIO(b"x"), 1, name, typ, kind)
        item, _ = kb.add_file(self.r, io.BytesIO(b"\x89PNG"), 4, "../../evil/截图.PNG", "ref", "visual")
        self.assertEqual((item["kind"], item["source"]), ("visual", f"refs/{item['id']}-截图.png"))
        with self.assertRaises(ValueError):  # the client sent fewer bytes than announced
            kb.add_file(self.r, io.BytesIO(b"abc"), 10, "a.md", "file")
        self.assertEqual(len(kb.load(self.r)), 1)

    def test_broken_office_file_is_kept_with_a_note(self):
        item, _ = kb.add_file(self.r, io.BytesIO(b"not a zip"), 9, "坏.docx", "file")
        self.assertIn("文字抽取失败", item["note"])
        self.assertNotIn("text", item)

    def test_add_url_and_deck(self):
        for url, typ, kind in (("ftp://x.com/a", "web", None), ("javascript:alert(1)", "web", None), ("https://x.com", "ref", "visual")):
            with self.assertRaises(ValueError):
                kb.add_url(self.r, url, typ, kind)
        a, c1 = kb.add_url(self.r, " https://example.com/r ", "web")
        b, c2 = kb.add_url(self.r, "https://example.com/r", "web")
        self.assertEqual((a["source"], c1, b["id"], c2), ("https://example.com/r", True, a["id"], False))
        cite, _ = kb.add_url(self.r, "https://example.com/r", "ref", "cite")  # same URL as a citation is a different entry
        self.assertNotEqual(cite["id"], a["id"])
        old = self.proj / "decks" / "old"
        old.mkdir()
        (old / "deck.html").write_text("<html></html>", encoding="utf-8")
        d, _ = kb.add_deck(self.r, old)
        self.assertEqual((d["type"], d["kind"], d["title"], d["source"]), ("ref", "deck", "old", str(old.resolve())))
        with self.assertRaises(ValueError):
            kb.add_deck(self.r, self.proj / "decks" / "nope")

    def test_update_and_remove_to_trash(self):
        item, _ = kb.add_file(self.r, io.BytesIO(b"# hi"), 4, "a.md", "file")
        (self.r / "cards").mkdir()
        (self.r / "cards" / f"{item['id']}.md").write_text("card", encoding="utf-8")
        self.assertEqual(kb.update(self.r, item["id"], status="ready", summary="s")["status"], "ready")
        kb.remove(self.r, item["id"])
        self.assertEqual(kb.load(self.r), [])
        trash = self.r / ".trash" / item["id"]
        self.assertEqual(sorted(p.name for p in trash.iterdir()), sorted([f"{item['id']}-a.md", f"{item['id']}.md", "item.json"]))
        with self.assertRaises(ValueError):
            kb.update(self.r, item["id"], status="ready")

    def test_remove_never_leaves_the_kb(self):
        outside = self.tmp / "keep.txt"
        outside.write_text("mine", encoding="utf-8")
        kb.save(self.r, [{"id": "k-aaaaaa", "type": "file", "title": "x", "source": "../../keep.txt", "status": "ready"}])
        kb.remove(self.r, "k-aaaaaa")
        self.assertEqual(outside.read_text(encoding="utf-8"), "mine")

    def test_save_is_atomic(self):
        kb.save(self.r, [{"id": "k-aaaaaa"}])
        kb.save(self.r, [{"id": "k-bbbbbb"}])
        self.assertEqual(kb.load(self.r), [{"id": "k-bbbbbb"}])
        self.assertEqual([p.name for p in self.r.iterdir()], ["index.json"])  # no temp files left behind

    def test_bad_index_raises_and_state_reports_it(self):
        g = kb.root(self.deck, "global")
        g.mkdir(parents=True)
        (g / "index.json").write_text("{oops", encoding="utf-8")
        with self.assertRaises(kb.BadIndex):
            kb.load(g)
        st = kb.state(self.deck)
        self.assertEqual(st["global"]["items"], [])
        self.assertTrue(st["global"]["error"])
        self.assertIsNone(st["project"]["error"])
        self.assertEqual((g / "index.json").read_text(encoding="utf-8"), "{oops")

    def test_skills_selection_and_state(self):
        sk = self.proj / ".claude" / "skills"
        (sk / "brand").mkdir(parents=True)
        (sk / "brand" / "SKILL.md").write_text('---\nname: brand-guide\ndescription: "公司品牌规范"\n---\n规则', encoding="utf-8")
        (sk / "plain").mkdir()
        (sk / "plain" / "SKILL.md").write_text("没有 frontmatter", encoding="utf-8")
        g = self.home / ".claude" / "skills" / "consult"
        g.mkdir(parents=True)
        (g / "SKILL.md").write_text("---\nname: consult\ndescription: 咨询式汇报\n---\n", encoding="utf-8")
        self.assertEqual([(s["name"], s["description"], s["scope"]) for s in kb.skills(self.deck)],
                         [("brand-guide", "公司品牌规范", "project"), ("plain", "", "project"), ("consult", "咨询式汇报", "global")])
        self.assertEqual(kb.selection(self.deck), [])
        kb.set_selection(self.deck, ["brand-guide"])
        self.assertEqual(kb.selection(self.deck), ["brand-guide"])
        for bad in ("brand-guide", [1], ["x" * 101]):
            with self.assertRaises(ValueError):
                kb.set_selection(self.deck, bad)
        (self.proj / "decks" / "old").mkdir()
        (self.proj / "decks" / "old" / "deck.html").write_text("x", encoding="utf-8")
        st = kb.state(self.deck)
        self.assertEqual((st["decks"], st["enabled"], len(st["skills"])), (["old"], ["brand-guide"], 3))

    def test_upload_does_not_lose_a_concurrent_update(self):  # review I1: --kb-done landing while a file streams in
        first, _ = kb.add_file(self.r, io.BytesIO(b"# a"), 3, "a.md", "file")
        r = self.r

        class Racing(io.BytesIO):
            def read(s, n=-1):
                if s.tell() == 0:
                    kb.update(r, first["id"], status="ready")
                return super().read(n)
        kb.add_file(self.r, Racing(b"# b"), 3, "b.md", "file")
        self.assertEqual([i["status"] for i in kb.load(self.r)], ["ready", "pending"])

    def test_save_retries_a_busy_index_and_cleans_up(self):  # review I1: Windows refuses to replace an open file
        real, calls = os.replace, []

        def flaky(a, b):
            calls.append(1)
            if len(calls) < 3:
                raise PermissionError("busy")
            return real(a, b)
        with mock.patch("kb.os.replace", flaky):
            kb.save(self.r, [{"id": "k-aaaaaa"}])
        self.assertEqual(kb.load(self.r), [{"id": "k-aaaaaa"}])
        with mock.patch("kb.os.replace", side_effect=PermissionError("busy")):
            with self.assertRaises(PermissionError):
                kb.save(self.r, [])
        self.assertEqual([p.name for p in self.r.iterdir()], ["index.json"])

    def test_extraction_never_breaks_an_upload(self):  # review I3
        bad_xlsx = zip_bytes({"xl/workbook.xml": f'<workbook {S_NS}><sheets><sheet name="s"/></sheets></workbook>',
                              "xl/worksheets/sheet1.xml": f'<worksheet {S_NS}><sheetData><row><c t="s"><v/></c></row></sheetData></worksheet>'})
        item, _ = kb.add_file(self.r, io.BytesIO(bad_xlsx), len(bad_xlsx), "坏表.xlsx", "file")
        self.assertIn("文字抽取失败", item["note"])
        with mock.patch("kb.extract", side_effect=RuntimeError("encrypted")):
            item2, _ = kb.add_file(self.r, io.BytesIO(DOCX), len(DOCX), "加密.docx", "file")
        self.assertIn("encrypted", item2["note"])
        with mock.patch("kb.save", side_effect=OSError("disk full")):
            with self.assertRaises(OSError):
                kb.add_file(self.r, io.BytesIO(b"# c"), 3, "c.md", "file")
        self.assertEqual(sorted(p.name.split("-", 2)[-1] for p in (self.r / "files").glob("k-*")), sorted(["坏表.xlsx", "加密.docx"]))

    def test_extracted_text_is_capped(self):  # review I4: a huge sheet must not be held whole in memory or on disk
        with mock.patch.object(kb, "MAX_TEXT_OUT", 10):
            item, _ = kb.add_file(self.r, io.BytesIO(DOCX), len(DOCX), "长.docx", "file")
        text = (self.r / item["text"]).read_text(encoding="utf-8")
        self.assertTrue(text.endswith("（已截断）"), text)
        self.assertLess(len(text), 40)

    def test_ids_from_the_index_are_checked(self):  # review M1: a hand-edited id must not reach the filesystem
        kb.save(self.r, [{"id": "../../evil", "type": "file", "title": "x", "source": "files/x.md", "status": "ready"}])
        with self.assertRaises(ValueError):
            kb.remove(self.r, "../../evil")
        self.assertFalse((self.tmp / "evil").exists())

    def test_digest_event_only_points_inside(self):  # review I2: a hand-edited index must not send Claude elsewhere
        outside = self.tmp / "secret.txt"
        outside.write_text("secret", encoding="utf-8")
        elsewhere = self.tmp / "other" / "decks" / "x"
        elsewhere.mkdir(parents=True)
        (elsewhere / "deck.html").write_text("x", encoding="utf-8")
        for item in ({"id": "k-aaaaaa", "type": "file", "title": "t", "source": str(outside)},
                     {"id": "k-aaaaaa", "type": "file", "title": "t", "source": "../../secret.txt"},
                     {"id": "k-aaaaaa", "type": "file", "title": "t", "source": "files/a.md", "text": str(outside)},
                     {"id": "k-aaaaaa", "type": "ref", "kind": "deck", "title": "t", "source": str(elsewhere)},
                     {"id": "k-aaaaaa", "type": "web", "title": "t", "source": "file:///C:/secret.txt"},
                     {"id": "k-aaaaaa", "type": "file", "title": "t"}):
            with self.assertRaises(ValueError, msg=item):
                kb.digest_event(self.deck, "project", item)

    def test_digest_event_and_card(self):
        item, _ = kb.add_file(self.r, io.BytesIO(PPTX), len(PPTX), "旧稿.pptx", "ref", "deck")
        ev = kb.digest_event(self.deck, "project", item)
        self.assertEqual((ev["type"], ev["kind"], ev["scope"], ev["kb_id"], ev["text"]), ("kb", "digest", "project", item["id"], "消化《旧稿》"))
        self.assertEqual(ev["source"], str(self.r / item["source"]))
        self.assertEqual(ev["extracted"], str(self.r / item["text"]))
        self.assertEqual(ev["card"], str(self.r / "cards" / f"{item['id']}.md"))
        web, _ = kb.add_url(self.r, "https://example.com", "web")
        wev = kb.digest_event(self.deck, "project", web)
        self.assertEqual((wev["source"], wev["snapshot"]), ("https://example.com", str(self.r / "web" / f"{web['id']}.md")))
        (self.r / "cards").mkdir(exist_ok=True)
        (self.r / "cards" / f"{item['id']}.md").write_text("## 摘要", encoding="utf-8")
        self.assertEqual(kb.card(self.deck, "project", item["id"]), "## 摘要")
        with self.assertRaises(ValueError):
            kb.card(self.deck, "project", "../index")


def post_raw(port, path, data, ctype="application/octet-stream", method="POST"):
    c = http.client.HTTPConnection("127.0.0.1", port, timeout=5)
    c.request(method, path, data, {"Host": f"127.0.0.1:{port}", "Content-Type": ctype})
    r = c.getresponse()
    out = (r.status, r.read())
    c.close()
    return out


class KbServerTests(ServerCase):
    def setUp(self):
        super().setUp()
        self.proj, self.home = self.tmp / "proj", self.tmp / "home"
        env = mock.patch.dict(os.environ, {"PPT_PROJECT_DIR": str(self.proj), "PPT_HOME_DIR": str(self.home)})
        env.start()
        self.addCleanup(env.stop)
        self.r = self.proj / "kb"

    def up(self, data, name, typ="file", kind="", scope="project", token="tok"):
        q = f"/api/kb/file?t={token}&scope={scope}&type={typ}&kind={kind}&name={quote(name)}"
        status, body = post_raw(self.port, q, data)
        return status, json.loads(body) if body.startswith(b"{") else body

    def test_upload_streams_and_asks_claude_to_digest(self):
        status, body = self.up(DOCX, "报告.docx")
        self.assertEqual((status, body["created"]), (200, True))
        item = body["item"]
        self.assertEqual(kb.load(self.r)[0]["id"], item["id"])
        ev = inbox(self.deck)[-1]
        self.assertEqual((ev["type"], ev["kind"], ev["kb_id"], ev["extracted"]), ("kb", "digest", item["id"], str(self.r / item["text"])))
        self.assertTrue(server.summary(ev).startswith(f"{ev['id']} kb kind:digest"))
        status, body = self.up(DOCX, "again.docx")
        self.assertEqual((status, body["created"]), (200, False))
        self.assertEqual(len([e for e in inbox(self.deck) if e["type"] == "kb"]), 1)  # no second digest for a duplicate

    def test_upload_rejections_reach_the_client(self):
        self.assertEqual(self.up(b"x" * 1000, "a.exe")[0], 400)
        self.assertEqual(self.up(b"x" * 1000, "a.png", typ="ref")[0], 400)
        self.assertEqual(self.up(b"x" * 1000, "a.md", token="bad")[0], 403)
        with mock.patch.object(kb, "MAX_FILE", 16):
            self.assertEqual(self.up(b"x" * 1000, "a.md")[0], 413)
        self.assertFalse(self.r.exists())

    def test_add_remove_redigest(self):
        def post(path, body):
            s, b = post_raw(self.port, f"{path}?t=tok", json.dumps(body).encode(), "application/json")
            return s, json.loads(b)
        self.assertEqual(post("/api/kb/add", {"scope": "project", "type": "web", "url": "ftp://x"})[0], 400)
        s, b = post("/api/kb/add", {"scope": "global", "type": "web", "url": "https://example.com"})
        self.assertEqual((s, b["created"]), (200, True))
        g = self.home / ".ppt-maker" / "kb"
        self.assertEqual(kb.load(g)[0]["source"], "https://example.com")
        (self.proj / "decks" / "old").mkdir(parents=True)
        (self.proj / "decks" / "old" / "deck.html").write_text("x", encoding="utf-8")
        s, b = post("/api/kb/add", {"scope": "project", "deck": "old"})
        self.assertEqual((s, b["item"]["kind"]), (200, "deck"))
        did = b["item"]["id"]
        kb.update(self.r, did, status="failed", note="打不开")
        s, b = post("/api/kb/redigest", {"scope": "project", "id": did})
        self.assertEqual((s, b["item"]["status"], b["item"]["note"]), (200, "pending", ""))
        self.assertEqual(inbox(self.deck)[-1]["kb_id"], did)
        s, b = post("/api/kb/remove", {"scope": "project", "id": did})
        self.assertEqual((s, kb.load(self.r)), (200, []))
        self.assertEqual(post("/api/kb/remove", {"scope": "project", "id": did})[0], 400)
        self.assertEqual(post("/api/kb/add", {"scope": "nowhere", "type": "web", "url": "https://x.com"})[0], 400)

    def test_bad_index_is_409_and_untouched(self):
        g = self.home / ".ppt-maker" / "kb"
        g.mkdir(parents=True)
        (g / "index.json").write_text("{oops", encoding="utf-8")
        st = json.loads(self.req("GET", "/api/kb?t=tok")[1])
        self.assertTrue(st["global"]["error"])
        s, _ = post_raw(self.port, "/api/kb/add?t=tok", json.dumps({"scope": "global", "type": "web", "url": "https://x.com"}).encode(), "application/json")
        self.assertEqual(s, 409)
        self.assertEqual(self.up(b"# x", "a.md", scope="global")[0], 409)
        self.assertEqual((g / "index.json").read_text(encoding="utf-8"), "{oops")

    def test_skills_put_and_skill_request(self):
        s, b = post_raw(self.port, "/api/kb/skills?t=tok", json.dumps({"skills": ["brand"]}).encode(), "application/json", "PUT")
        self.assertEqual((s, json.loads(b)["enabled"]), (200, ["brand"]))
        self.assertEqual(kb.selection(self.deck), ["brand"])
        self.assertEqual(post_raw(self.port, "/api/kb/skills?t=tok", b'{"skills": "brand"}', "application/json", "PUT")[0], 400)
        s, _ = post_raw(self.port, "/api/kb/skill?t=tok", json.dumps({"scope": "global", "text": "品牌规范：主色蓝"}).encode(), "application/json")
        ev = inbox(self.deck)[-1]
        self.assertEqual((s, ev["kind"], ev["dir"], ev["text"]), (200, "skill", str(self.home / ".claude" / "skills"), "品牌规范：主色蓝"))
        self.assertEqual(post_raw(self.port, "/api/kb/skill?t=tok", b'{"scope": "project", "text": " "}', "application/json")[0], 400)

    def test_state_and_card(self):
        st = json.loads(self.req("GET", "/api/kb?t=tok")[1])
        self.assertEqual((st["project"]["items"], st["skills"], st["enabled"]), ([], [], []))
        self.assertEqual(self.req("GET", "/api/kb")[0], 403)
        _, body = self.up(b"# x", "a.md")
        (self.r / "cards").mkdir()
        (self.r / "cards" / f"{body['item']['id']}.md").write_text("## 摘要\n要点", encoding="utf-8")
        s, b = self.req("GET", f"/api/kb/card?t=tok&scope=project&id={body['item']['id']}")
        self.assertEqual((s, json.loads(b)["text"]), (200, "## 摘要\n要点"))
        self.assertEqual(self.req("GET", "/api/kb/card?t=tok&scope=project&id=k-zzzzzz")[0], 404)

    def test_kb_done_and_fail_cli(self):
        _, body = self.up(b"# x", "a.md")
        kid = body["item"]["id"]
        with contextlib.redirect_stdout(io.StringIO()):  # keep the CLI's one-line confirmation out of the test report
            server.main(["--kb-done", str(self.deck), "project", kid, "年报", "营收", "12.3 亿"])
        self.assertEqual({k: kb.load(self.r)[0][k] for k in ("status", "title", "summary")}, {"status": "ready", "title": "年报", "summary": "营收 12.3 亿"})
        with contextlib.redirect_stdout(io.StringIO()):
            server.main(["--kb-fail", str(self.deck), "project", kid, "加密", "PDF"])
        self.assertEqual((kb.load(self.r)[0]["status"], kb.load(self.r)[0]["note"]), ("failed", "加密 PDF"))

    def test_intent_kb_field(self):
        self.assertEqual(self.req("POST", "/api/event?t=tok", {"type": "intent", "text": "x", "kb": ["k-abc123"]})[0], 200)
        e = inbox(self.deck)[-1]
        self.assertEqual(e["kb"], ["k-abc123"])
        self.assertIn("kb:k-abc123", server.summary(e))
        for bad in ([5], ["k" * 21], ["k-a"] * 21):
            self.assertEqual(self.req("POST", "/api/event?t=tok", {"type": "intent", "text": "x", "kb": bad})[0], 400)

    def test_a_slow_upload_does_not_block_other_requests(self):  # review I4: the stream is read outside LOCK
        s = socket.create_connection(("127.0.0.1", self.port))
        s.sendall(f"POST /api/kb/file?t=tok&scope=project&type=file&name=a.md HTTP/1.1\r\nHost: 127.0.0.1:{self.port}\r\n"
                  f"Content-Length: 100\r\n\r\n# partial".encode())
        _time.sleep(0.3)
        try:
            c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=3)
            c.request("GET", "/api/kb?t=tok", headers={"Host": f"127.0.0.1:{self.port}"})
            self.assertEqual(c.getresponse().status, 200)
            c.close()
        finally:
            s.close()

    def test_busy_index_is_503_not_a_dropped_connection(self):  # review I1
        with mock.patch("kb.save", side_effect=PermissionError("busy")):
            s, b = post_raw(self.port, "/api/kb/add?t=tok", json.dumps({"scope": "project", "type": "web", "url": "https://x.com"}).encode(), "application/json")
        self.assertEqual(s, 503)

    def test_redigest_refuses_paths_outside_the_kb(self):  # review I2
        kb.save(self.r, [{"id": "k-aaaaaa", "type": "file", "title": "t", "source": str(self.tmp / "deck" / "deck.html"), "status": "failed"}])
        s, _ = post_raw(self.port, "/api/kb/redigest?t=tok", json.dumps({"scope": "project", "id": "k-aaaaaa"}).encode(), "application/json")
        self.assertEqual(s, 400)
        self.assertFalse(any(e["type"] == "kb" for e in inbox(self.deck)))

    def test_kb_add_cli(self):  # so Claude never edits index.json by hand
        src = self.tmp / "笔记.md"
        src.write_text("# 笔记", encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()) as out:
            server.main(["--kb-add", str(self.deck), "project", "file", str(src)])
            server.main(["--kb-add", str(self.deck), "global", "web", "https://example.com"])
        items = kb.load(self.r)
        self.assertEqual((items[0]["title"], items[0]["status"]), ("笔记", "pending"))
        self.assertEqual(kb.load(self.home / ".ppt-maker" / "kb")[0]["source"], "https://example.com")
        self.assertIn(items[0]["id"], out.getvalue())  # Claude needs the id (and the card path) to digest it
        self.assertIn("card", out.getvalue())

    def test_tail_catches_up_kb_events(self):
        self.up(b"# x", "a.md")
        lines, done = [], threading.Event()
        th = threading.Thread(target=server.tail, args=(self.deck, lines.append, done.is_set, 0.05), daemon=True)
        th.start()
        _time.sleep(0.3)
        done.set()
        th.join(2)
        self.assertTrue(any(" kb kind:digest " in l for l in lines), lines)


class BatchTests(ServerCase):
    def item(self, text, slide_id="s-cover", **kw):
        return {"slide": 1, "slide_id": slide_id, "rect": {"x": 0, "y": 0, "w": 100, "h": 50}, "cells": "1-2/1-1", "blocks": ["b-cover-title"], "text": text, **kw}

    def post(self, ev):
        return self.req("POST", "/api/event?t=tok", ev)[0]

    def test_batch_items_saved_with_their_shots(self):
        self.assertEqual(self.post({"type": "intent", "kind": "batch", "text": "整页改成对比结构",
                                    "items": [self.item("标题改短", shot=PNG), self.item("画一个流程", kind="sketch", notes=["开始"])]}), 200)
        e = inbox(self.deck)[-1]
        self.assertEqual([i["text"] for i in e["items"]], ["标题改短", "画一个流程"])
        self.assertEqual(e["items"][0]["shot"], f".ppt/shots/{e['id']}-1.png")
        self.assertTrue((self.deck / e["items"][0]["shot"]).read_bytes().startswith(b"\x89PNG"))
        self.assertNotIn("shot", e["items"][1])
        self.assertIn(" kind:batch ", server.summary(e))
        self.assertIn(" items:2 ", server.summary(e))

    def test_batch_validation(self):
        ok = self.item("x")
        for bad in ({"type": "intent", "kind": "batch", "text": "x"},
                    {"type": "intent", "kind": "batch", "items": "x"},
                    {"type": "intent", "kind": "batch", "items": []},
                    {"type": "intent", "kind": "batch", "items": [ok] * 31},
                    {"type": "intent", "kind": "batch", "items": [{**ok, "type": "intent"}]},
                    {"type": "intent", "kind": "batch", "items": [{**ok, "kind": "draw-icon"}]},
                    {"type": "intent", "kind": "batch", "items": [{**ok, "blocks": "x"}]},
                    {"type": "intent", "kind": "batch", "items": [{**ok, "items": [ok]}]},
                    {"type": "intent", "text": "x", "items": [ok]}):
            self.assertEqual(self.post(bad), 400, bad)

    def test_batch_fields_are_strict(self):  # review M5: nothing extra rides inside a batch, and only intents batch
        ok = self.item("x")
        for bad in ({"type": "answer", "re": "e1", "kind": "batch", "items": [ok]},
                    {"type": "lint", "kind": "batch", "items": [ok]},
                    {"type": "intent", "kind": "batch", "items": [{**ok, "re": "e9"}]},
                    {"type": "intent", "kind": "batch", "items": [{**ok, "issues": [{"kind": "x"}]}]}):
            self.assertEqual(self.post(bad), 400, bad)

    def test_batch_reply_follows_an_item_still_on_the_deck(self):  # review I3: item 1's slide may be gone
        self.post({"type": "intent", "kind": "batch", "items": [self.item("旧页上的", slide_id="s-gone"), self.item("封面上的")]})
        self.post({"type": "answer", "re": "e1", "text": "再改改"})
        server.reply(self.deck, "e1", "改好了")
        ev = {e["id"]: e for e in json.loads(self.req("GET", "/api/state?t=tok")[1])["events"]}
        self.assertEqual((ev["e1-1"]["reply"], ev["e1-2"]["reply"], ev["e2"]["re"]), (None, "改好了", "e1-2"))

    def test_state_expands_batches(self):
        self.post({"type": "intent", "kind": "batch", "items": [self.item("一"), self.item("二", slide_id="s-2")]})
        self.post({"type": "answer", "re": "e1", "text": "第二条再大一点"})
        server.reply(self.deck, "e1", "两处都改好了")
        ev = {e["id"]: e for e in json.loads(self.req("GET", "/api/state?t=tok")[1])["events"]}
        self.assertEqual((ev["e1-1"]["parent"], ev["e1-1"]["reply"], ev["e1-1"]["status"]), ("e1", "两处都改好了", "done"))
        self.assertEqual((ev["e1-2"]["slide_id"], ev["e1-2"]["reply"]), ("s-2", None))
        self.assertNotIn("e1", ev)
        self.assertEqual(ev["e2"]["re"], "e1-1")


REVEAL = '<html><body><div class="reveal"><div class="slides"><section><h1>导入测试</h1></section><section>二</section></div></div><script>window.x=1</script></body></html>'


class ImporterTests(ServerCase):
    def save(self, data, name):
        return importer.save(self.deck, io.BytesIO(data), len(data), name)

    def test_hint(self):
        self.assertEqual(importer.hint(TEMPLATE.read_text(encoding="utf-8")), "ppt-maker")
        self.assertEqual(importer.hint(REVEAL), "slides")
        self.assertEqual(importer.hint("<section>a</section>" * 3), "slides")
        self.assertEqual(importer.hint("<article><h1>报告</h1><p>正文</p></article>"), "page")

    def test_save_html_and_zip(self):
        iid, d, entry = self.save(REVEAL.encode(), "../旧稿.html")
        self.assertRegex(iid, r"^i-[0-9a-z]{6}$")
        self.assertEqual((d, entry.name, entry.read_text(encoding="utf-8")), (self.deck / "imports" / iid, "旧稿.html", REVEAL))
        z = zip_bytes({"deep/a/page.html": "<p>x</p>", "site/index.html": "<h1>首页</h1>", "site/img/a.png": b"\x89PNG", "site/about.html": "<p>y</p>"})
        _, d2, entry2 = self.save(z, "site.zip")
        self.assertEqual(entry2.relative_to(d2).as_posix(), "site/index.html")
        self.assertTrue((d2 / "site" / "img" / "a.png").is_file())
        self.assertFalse(any(p.suffix == ".zip" for p in d2.iterdir()))

    def test_unsafe_or_empty_zips_are_refused(self):
        for files in ({"../evil.html": "x"}, {"/abs.html": "x"}, {"C:/win.html": "x"}, {"a.txt": "no html here"}):
            with self.assertRaises(ValueError, msg=files):
                self.save(zip_bytes(files), "x.zip")
        with mock.patch.object(importer, "MAX_FILES", 2):
            with self.assertRaises(ValueError):
                self.save(zip_bytes({"a.html": "1", "b.html": "2", "c.html": "3"}), "x.zip")
        with self.assertRaises(ValueError):
            self.save(b"<p>x</p>", "x.exe")
        self.assertEqual(list((self.deck / "imports").iterdir()), [])  # nothing half-imported is left behind
        self.assertFalse((self.tmp / "evil.html").exists())

    def test_broken_zips_are_a_clear_400(self):  # review I2: not a dropped connection
        self.assertEqual(self.up(b"not a zip at all", "a.zip")[0], 400)
        self.assertEqual(self.up(zip_bytes({"a": "file", "a/b.html": "<p>x</p>"}), "clash.zip")[0], 400)
        self.assertEqual(list((self.deck / "imports").iterdir()), [])

    def test_a_member_named_like_the_upload_is_harmless(self):  # review M8
        _, d, entry = self.save(zip_bytes({".upload.zip": "x", "index.html": "<h1>x</h1>"}), "site.zip")
        self.assertEqual(entry.name, "index.html")
        self.assertEqual((d / ".upload.zip").read_text(), "x")

    def test_original_link_is_url_safe(self):  # review M10: '#' and spaces in a file name
        self.up(REVEAL.encode(), "旧稿 #2.html")
        e = inbox(self.deck)[-1]
        self.assertEqual(e["original"], f"/imports/{e['import_id']}/" + quote("旧稿 #2.html"))
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", e["original"], headers={"Host": f"127.0.0.1:{self.port}"})
        self.assertEqual(c.getresponse().status, 200)
        c.close()

    def test_every_deck_file_is_sandboxed_and_the_page_ignores_base(self):  # review I4: defence in depth
        (self.deck / "assets").mkdir()
        (self.deck / "assets" / "x.svg").write_text('<svg xmlns="http://www.w3.org/2000/svg"><script>alert(1)</script></svg>', encoding="utf-8")
        for path, want in (("/assets/x.svg", "sandbox"), ("/?t=tok", "base-uri 'none'")):
            c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
            c.request("GET", path, headers={"Host": f"127.0.0.1:{self.port}"})
            self.assertEqual(c.getresponse().getheader("Content-Security-Policy"), want, path)
            c.close()

    def test_a_declared_charset_further_down_is_respected(self):  # review M9
        gbk = ("<html><head><style>" + "a{}" * 3000 + '</style><meta charset="gbk"></head><body>中文</body></html>').encode("gbk")
        self.up(gbk, "gbk.html")
        e = inbox(self.deck)[-1]
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", e["original"], headers={"Host": f"127.0.0.1:{self.port}"})
        self.assertNotIn("utf-8", c.getresponse().getheader("Content-Type"))
        c.close()

    def test_zip_bomb_stops_at_the_real_size(self):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
            z.writestr("index.html", "<p>" + "x" * 5000 + "</p>")
        with mock.patch.object(importer, "MAX_IMPORT", 1000):
            with self.assertRaises(ValueError):
                self.save(buf.getvalue(), "bomb.zip")

    def up(self, data, name, mode="append", token="tok"):
        return post_raw(self.port, f"/api/import?t={token}&mode={mode}&name={quote(name)}", data, "text/html")

    def test_import_endpoint_and_event(self):
        s, b = self.up(REVEAL.encode(), "旧稿.html")
        self.assertEqual(s, 200)
        e = inbox(self.deck)[-1]
        self.assertEqual((e["id"], e["type"], e["kind"], e["mode"], e["hint"]), (json.loads(b)["id"], "import", "file", "append", "slides"))
        self.assertTrue(Path(e["source"]).is_file())
        self.assertEqual(e["original"], f"/imports/{e['import_id']}/" + quote("旧稿.html"))
        self.assertIn(" import kind:file ", server.summary(e))
        for args, code in (((b"x" * 1000, "a.html", "sideways"), 400), ((b"x" * 1000, "a.exe"), 400), ((b"x" * 1000, "a.html", "append", "bad"), 403)):
            self.assertEqual(self.up(*args)[0], code, args)
        with mock.patch.object(importer, "MAX_IMPORT", 16):
            self.assertEqual(self.up(b"x" * 1000, "a.html")[0], 413)

    def test_original_is_sandboxed(self):
        self.up(REVEAL.encode(), "旧稿.html")
        e = inbox(self.deck)[-1]
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", e["original"], headers={"Host": f"127.0.0.1:{self.port}"})
        r = c.getresponse()
        body = r.read().decode()
        c.close()
        self.assertEqual((r.status, r.getheader("Content-Security-Policy")), (200, "sandbox"))
        self.assertIn("charset=utf-8", r.getheader("Content-Type"))  # no <meta charset>: Chinese must not turn into mojibake
        self.assertIn("导入测试", body)
        c = http.client.HTTPConnection("127.0.0.1", self.port, timeout=5)
        c.request("GET", "/?t=tok", headers={"Host": f"127.0.0.1:{self.port}"})
        self.assertNotIn("sandbox", c.getresponse().getheader("Content-Security-Policy"))  # the editor page itself still runs
        c.close()

    def test_url_import_state_and_cli(self):
        def post(body):
            return post_raw(self.port, "/api/import/url?t=tok", json.dumps(body).encode(), "application/json")[0]
        self.assertEqual(post({"url": "ftp://x.com", "mode": "append"}), 400)
        self.assertEqual(post({"url": "https://example.com/a", "mode": "replace"}), 200)
        e = inbox(self.deck)[-1]
        self.assertEqual((e["type"], e["kind"], e["mode"], e["source"]), ("import", "url", "replace", "https://example.com/a"))
        server.reply(self.deck, e["id"], "导入了 3 页")
        st = {x["id"]: x for x in json.loads(self.req("GET", "/api/state?t=tok")[1])["events"]}
        self.assertEqual((st[e["id"]]["type"], st[e["id"]]["status"], st[e["id"]]["reply"]), ("import", "done", "导入了 3 页"))
        src = self.tmp / "旧.html"
        src.write_text(REVEAL, encoding="utf-8")
        with contextlib.redirect_stdout(io.StringIO()) as out:
            server.main(["--import", str(self.deck), str(src), "replace"])
        e2 = inbox(self.deck)[-1]
        self.assertEqual((e2["type"], e2["mode"], e2["hint"]), ("import", "replace", "slides"))
        self.assertIn(e2["id"], out.getvalue())
        self.assertEqual(e2["origin_dir"], str(src.resolve().parent))  # review M11: its pictures stay next to the original


class CliTests(ServerCase):
    def test_reply_drives_status(self):
        server.reply(self.deck, "e1", ack=True)
        self.assertEqual(server.statuses(self.deck), {"e1": "processing"})
        server.reply(self.deck, "e1", "做完了")
        server.reply(self.deck, "e1", ack=True)           # late ack must not undo done
        self.assertEqual(server.statuses(self.deck), {"e1": "done"})

    def test_summary_one_line(self):
        s = server.summary({"id": "e3", "type": "answer", "re": "e1", "slide_id": "s1", "blocks": ["a", "b"], "text": "第一行\n第二行" + "长" * 100})
        self.assertTrue(s.startswith("e3 answer re:e1 slide:s1 [a,b] 第一行 第二行"))
        self.assertNotIn("\n", s)
        self.assertIn("[x]", server.summary({"id": "e4", "type": "edit", "slide_id": "s1", "block": "x"}))

    def test_tail_catch_up_then_follow(self):
        for t in ("旧的已完成", "旧的处理中", "旧的排队"):
            self.req("POST", "/api/event?t=tok", {"type": "intent", "slide_id": "s-cover", "text": t})
        server.reply(self.deck, "e1", "ok")
        server.reply(self.deck, "e2", ack=True)
        lines, done = [], threading.Event()
        th = threading.Thread(target=server.tail, args=(self.deck, lines.append, done.is_set, 0.05), daemon=True)
        th.start()
        _time.sleep(0.3)
        self.assertEqual([l.split()[1] if l.startswith("最近完成") else l.split()[0] for l in lines], ["e1", "e2", "e3"])  # finished work first, as context
        self.req("POST", "/api/event?t=tok", {"type": "lint", "issues": [{"block": "b", "kind": "overflow"}]})
        _time.sleep(0.3)
        done.set()
        th.join(2)
        self.assertEqual([l.split()[0] for l in lines][1:], ["e2", "e3", "e4"])
        self.assertTrue((self.deck / ".ppt" / "heartbeat").exists())

    def test_cli_reply(self):
        server.main(["--reply", str(self.deck), "e9", "已", "完成", "--ack"])
        server.main(["--reply", str(self.deck), "e9", "已", "完成"])
        rows = server.read_jsonl(self.deck / ".ppt" / "outbox.jsonl")
        self.assertEqual([(r["re"], r["text"], r["ack"]) for r in rows], [("e9", "已 完成", True), ("e9", "已 完成", False)])


if __name__ == "__main__":
    unittest.main()
